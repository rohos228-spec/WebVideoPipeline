"""Книгу на боевом пути SaaS не читают. Это требование стоит проверкой.

Запись в книгу в SaaS выключена (`XLSX_WRITE`), значит файла на диске нет.
Любой `load_workbook` по несуществующему пути — это `FileNotFoundError`
посреди шага, за который клиент уже заплатил.

Читающих книгу мест двадцать три, и почти все из них — операторский
инструментарий: версионирование, импорт, правка ячейки, лист массового
проекта. В SaaS они просто не запускаются. Опасны те, что лежат на пути
генерации, и главный из них один: подбор референсов кадра.

Он написан правильно — сначала база, книга запасным путём, — но «написан
правильно» это состояние, а не гарантия. Здесь оно становится гарантией.
"""

from __future__ import annotations

import pytest
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models import Artifact, ArtifactKind, Base, Entity, Frame, FrameStatus, Project, ProjectStatus
from app.settings import settings


@pytest.fixture
def forbid_workbook(monkeypatch):
    """Любое чтение книги в этом тесте — падение с внятным текстом."""
    import openpyxl

    def _boom(*args, **kwargs):
        raise AssertionError(f"боевой путь прочитал книгу: {args[:1]}")

    monkeypatch.setattr(openpyxl, "load_workbook", _boom)
    # Модули импортируют функцию по имени — подменяем и у них.
    import app.orchestrator.steps.generate_images as gi

    monkeypatch.setattr(gi, "load_workbook", _boom, raising=False)
    return _boom


@pytest.fixture
async def project(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data")
    monkeypatch.setattr(settings, "xlsx_write", False)
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'x.db'}", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)

    async with factory() as s:
        p = Project(slug="saas-path", topic="без книги", status=ProjectStatus.image_prompts_ready)
        s.add(p)
        await s.flush()
        s.add(
            Entity(
                project_id=p.id,
                type="character",
                code="c01",
                name="Герой",
                attrs={"look": "мужчина", "clothes": "куртка"},
            )
        )
        s.add(
            Frame(
                project_id=p.id,
                number=1,
                voiceover_text="текст",
                image_prompt="Фон: вагон\nДействие: стоит",
                status=FrameStatus.image_prompt_ready,
                attrs={"characters": "c01", "shot01_bg": "вагон"},
            )
        )
        await s.commit()
    yield factory, engine
    await engine.dispose()


async def test_no_workbook_on_disk_in_saas(project) -> None:
    """Основа всей проверки: файла действительно нет."""
    factory, _ = project
    async with factory() as s:
        p = await s.get(Project, 1)
        assert not (p.data_dir / "project.xlsx").exists()


async def test_refs_come_from_the_database_without_touching_the_workbook(project, forbid_workbook) -> None:
    """Референсы кадра берутся из базы, а книга не открывается вовсе.

    Это главное место, где боевой путь мог бы упереться в отсутствующий файл:
    подбор персонажей и предметов кадра. Порядок «сначала база» здесь уже
    был; тест превращает его из состояния кода в требование.
    """
    from app.orchestrator.steps.generate_images import _load_refs_for_frame

    factory, _ = project
    async with factory() as s:
        p = await s.get(Project, 1)
        # Реф лежит в базе — книга не понадобится даже как запасной путь.
        chars = p.data_dir / "characters"
        chars.mkdir(parents=True, exist_ok=True)
        (chars / "c01.png").write_bytes(b"png")
        s.add(
            Artifact(
                project_id=p.id,
                kind=ArtifactKind.hero_reference,
                uuid="ref-c01",
                path=str(chars / "c01.png"),
                meta={"code": "c01"},
            )
        )
        await s.commit()

    async with factory() as s:
        p = await s.get(Project, 1)
        refs = await _load_refs_for_frame(s, p, 1)
    assert isinstance(refs, list)


async def test_missing_workbook_is_not_a_crash(project) -> None:
    """Даже когда в базе рефов нет, отсутствие книги не роняет шаг.

    Запасной путь обязан упираться в «файла нет» и молча возвращать пусто:
    кадр без референса рисуется по описанию, а упавший шаг не рисуется вовсе.
    """
    from app.orchestrator.steps.generate_images import _load_refs_for_frame

    factory, _ = project
    async with factory() as s:
        p = await s.get(Project, 1)
        refs = await _load_refs_for_frame(s, p, 1)
    assert refs == []


async def test_prompt_canon_needs_no_workbook(project, forbid_workbook) -> None:
    """Канон внешности и фона собирается из базы, а не из книги.

    Он встал на боевой путь последним, и это тот самый случай, когда новая
    правильная вещь могла бы притащить за собой чтение файла.
    """
    from app.services.img_pr_canon import enforce_canon_in_prompts

    factory, _ = project
    async with factory() as s:
        p = await s.get(Project, 1)
        stats = await enforce_canon_in_prompts(s, p)
    assert stats["frames"] == 1
