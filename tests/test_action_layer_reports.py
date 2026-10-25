from __future__ import annotations

import itertools
import json
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.models import Artifact, ArtifactKind, Entity, Frame, Project, ProjectStatus
from app.orchestrator.agents.summarizer import download_links, known_text
from app.services import prompt_library as pl
from app.services import spend_guard
from app.services.action_layer import builtin_insight as bi
from app.services.action_layer import report_gpt as rg
from app.services.action_layer import report_template as tpl
from app.services.action_layer.context import ActionCallError
from app.services.action_layer.planner_catalog import select_catalog
from app.services.action_layer.planner_rules import notes_for
from app.services.action_layer.spec import ActionRegistry, ensure_builtin_actions, registry
from app.settings import settings
from tests.action_layer_support import env  # noqa: F401

_SLUGS = itertools.count(1)


class FakeCtx:
    def __init__(self, routes):
        self.routes = routes

    async def call(self, method, path, *, params=None, json=None, timeout=0):
        val = self.routes.get((method, path))
        if val is None:
            raise ActionCallError(404, "нет")
        return val


@pytest.fixture
def specs():
    reg = ActionRegistry()
    bi.register_insight(reg)
    return reg


@pytest.fixture
def prompts(monkeypatch):
    monkeypatch.setattr(pl, "excel_gpt_prompt_exists", lambda name: name in ("razmetka", "geroi"))
    monkeypatch.setattr(pl, "load_file_meta", lambda step: {"razmetka": {"title": "Разметка кадров"},
                                                           "geroi": {"title": "Персонажи"}})


def _meta():
    return {
        "canvas_graph": {"nodes": [
            {"id": "n_plan", "type": "plan", "data": {"label": "План"}},
            {"id": "excel_gpt_1", "type": "excel_gpt", "data": {"label": "Работа с GPT", "slotIndex": 1}},
            {"id": "excel_gpt_2", "type": "excel_gpt", "data": {"label": "Работа с GPT", "slotIndex": 2}},
            {"id": "n_img", "type": "images", "data": {"label": "Картинки"}},
        ], "edges": []},
        "prompt_slot_variants": {"excel_gpt_1": {"main": "razmetka"}, "excel_gpt_2": {"main": "geroi"}},
        "scene_registry": [{"id_scene": "s1"}, {"id_scene": "s2"}, {"id_scene": "old", "активен": False}],
    }


REPLY = """Готово.
```json
{"ops": [
  {"target": "frame", "frame_uuid": "uuid-1", "fields": {"промт_картинки": "кот", "место": "кухня"}},
  {"target": "frame", "frame_uuid": "2", "fields": {"image_prompt": "пёс"}},
  {"target": "frame", "frame_uuid": "uuid-1", "fields": {"main_action": "прыгает"}},
  {"target": "project", "fields": {"general_plan": "план"}}
],
 "characters": [{"id": "c01", "имя": "Кот Барсик"}],
 "scenes": [{"id_scene": "s1", "место": "кухня"}, {"id_scene": "s2", "место": "двор"}]}
```"""


async def _seed(env, *, title="Ролик про кота", meta=None, frames=3, images=0, videos=0, usd=None,
                status="images_ready"):
    async with env.factory() as s:
        p = Project(topic="тема", title=title, slug=f"rep-{next(_SLUGS)}",
                    status=ProjectStatus(status), meta=dict(meta if meta is not None else _meta()))
        s.add(p)
        await s.flush()
        for i in range(1, frames + 1):
            f = Frame(project_id=p.id, number=i, voiceover_text=f"р{i}", attrs={}, uuid=f"uuid-{i}")
            s.add(f)
        await s.flush()
        for i in range(images):
            s.add(Artifact(project_id=p.id, kind=ArtifactKind.scene_image, uuid=f"im-{p.id}-{i}", path="x.png"))
        for i in range(videos):
            s.add(Artifact(project_id=p.id, kind=ArtifactKind.scene_video, uuid=f"vd-{p.id}-{i}", path="x.mp4"))
        s.add(Entity(project_id=p.id, type="character", code="c01", name="Барсик"))
        if usd is not None:
            s.add(spend_guard.SpendLedger(created_at=datetime.utcnow(), day=spend_guard._today(), kind="actual",
                                          provider="kie", project_id=p.id, run_id="", source="video", usd=usd,
                                          credits=0.0, note="", request_id="", ref=""))
        await s.commit()
        return p.id, p.slug


def _write_reply(slug, key, text):
    up = settings.data_dir / "videos" / slug / "excel_gpt_uploads" / key
    up.mkdir(parents=True, exist_ok=True)
    (up / "gpt_reply.txt").write_text(text, encoding="utf-8")


def _routes(pid, slug, meta, runs=None):
    detail = {"id": pid, "slug": slug, "title": "Ролик про кота", "topic": "тема", "status": "images_ready",
              "meta": meta, "prompt_overrides": {}}
    return {("GET", f"/projects/{pid}"): detail,
            ("GET", f"/projects/{pid}/run"): {"id": 1, "node_runs": runs or []}}


RUNS = [{"node_key": "excel_gpt_1", "node_type": "excel_gpt", "status": "done", "attempts": 1},
        {"node_key": "n_plan", "node_type": "plan", "status": "done", "attempts": 1}]


def test_writes_from_reply_maps_frames_fields_characters_scenes() -> None:
    w = rg.writes_from_reply(REPLY, {"uuid-1": 1, "uuid-2": 2})
    assert w["parsed"] and w["frames"] == {1: ["image_prompt", "place", "main_action"], 2: ["image_prompt"]}
    assert w["project_fields"] == ["general_plan"]
    assert w["characters"] == ["c01 — Кот Барсик"] and w["scenes"] == ["s1 — кухня", "s2 — двор"]
    assert rg.writes_from_reply("нет json", {})["parsed"] is False
    assert rg.writes_from_reply(None, {})["frames"] == {}


def test_read_reply_stays_inside_project_folder(tmp_path) -> None:
    (tmp_path / "excel_gpt_uploads" / "k").mkdir(parents=True)
    (tmp_path / "excel_gpt_uploads" / "k" / "gpt_reply.txt").write_text("ok", encoding="utf-8")
    (tmp_path / "gpt_reply.txt").write_text("secret", encoding="utf-8")
    assert rg.read_reply(tmp_path, "k") == "ok"
    assert rg.read_reply(tmp_path, "..") is None and rg.read_reply(tmp_path, "../..") is None
    assert rg.read_reply(tmp_path, "missing") is None


async def test_report_nodes_has_gpt_section_with_prompt_model_and_writes(env, specs, prompts) -> None:
    meta = _meta()
    pid, slug = await _seed(env, meta=meta)
    _write_reply(slug, "excel_gpt_1", REPLY)
    ctx = FakeCtx(_routes(pid, slug, meta, RUNS))
    report = specs.get("report.nodes").executor

    data = await report(ctx, {"project_id": pid, "kind": "json"})
    gpt = {g["node_key"]: g for g in data["gpt"]}
    assert list(gpt) == ["excel_gpt_1", "excel_gpt_2"]
    one, two = gpt["excel_gpt_1"], gpt["excel_gpt_2"]
    assert one["label"] == "Работа с GPT 1" and two["label"] == "Работа с GPT 2"
    assert one["prompt_title"] == "Разметка кадров (razmetka)" and two["prompt_title"] == "Персонажи (geroi)"
    assert one["model"] and one["status_ru"] == "готово" and two["status_ru"] == "не запускалась"
    assert one["frames_written"] == 2 and one["frames"] == {"1": ["image_prompt", "place", "main_action"],
                                                           "2": ["image_prompt"]}
    assert one["fields_ru"] == "промт картинки, место, главное действие"
    assert one["characters_count"] == 1 and one["scenes_count"] == 2 and "проект: общий план" in one["note"]
    assert two["frames_written"] == 0 and two["note"] == ""
    assert data["summary"]["gpt_nodes"] == 2
    assert data["summary"]["characters_in_project"] == 1 and data["summary"]["scenes_in_registry"] == 2
    assert data["text_ru"].splitlines()[1] == "Работа с GPT: нод 2, готово 1."
    assert "Работа с GPT 1: промт «Разметка кадров (razmetka)»" in data["text_ru"]
    assert "заполнила кадров 2" in data["text_ru"] and "сцен 2" in data["text_ru"]

    md = await report(ctx, {"project_id": pid, "kind": "md"})
    assert "### Работа с GPT" in md["markdown"] and "| Кадр |" in md["markdown"]

    chat = await report(ctx, {"project_id": pid})
    assert chat["format"] == "chat" and chat["file_format"] == "html"
    assert chat["download_url"] == f"/api/projects/{pid}/reports/{chat['file_name']}"
    assert chat["file_name"].endswith(".html") and chat["path"] == chat["html_path"]
    xlsx_url = f"/api/projects/{pid}/reports/{chat['xlsx_path'].replace(chr(92), '/').rsplit('/', 1)[-1]}"
    assert "### Работа с GPT" in chat["markdown"] and chat["download_url"] in chat["markdown"]
    first = chat["text_ru"].splitlines()[0]
    assert first.startswith("Отчёт по нодам") and f"({chat['download_url']})" in first and xlsx_url in first
    page = (await env.client.get(chat["download_url"])).text
    assert "Дневник прогонов нод" in page and "Кто по очереди пишет одно поле" in page
    assert "Ячейка 1" in page and "прыгает" in page and "временн" not in page
    wb = load_workbook(chat["xlsx_path"])
    assert wb.sheetnames == ["Ноды", tpl.GPT_SHEET, tpl.GPT_FRAMES_SHEET, tpl.GPT_CHARACTERS_SHEET, "Итог"]
    ws = wb[tpl.GPT_SHEET]
    assert [c.value for c in ws[1]] == [label for _k, label in tpl.GPT_NODE_COLUMNS]
    assert ws.max_row == 3 and ws["A2"].value == "Работа с GPT 1"
    frames = [[c.value for c in row] for row in wb[tpl.GPT_FRAMES_SHEET].iter_rows(min_row=2)]
    # Строка на поле, значение целиком (данные в ячейках не теряются).
    assert sorted(frames) == sorted([["Работа с GPT 1", "1", "промт картинки", "кот", None],
                                     ["Работа с GPT 1", "1", "место", "кухня", None],
                                     ["Работа с GPT 1", "1", "главное действие", "прыгает", None],
                                     ["Работа с GPT 1", "2", "промт картинки", "пёс", None]])
    got = await env.client.get(xlsx_url)
    assert got.status_code == 200 and got.content[:2] == b"PK"
    assert download_links(chat) == [f"[Скачать отчёт (html)]({chat['download_url']})",
                                    f"[Скачать отчёт (xlsx)]({xlsx_url})"]
    xl = await report(ctx, {"project_id": pid, "kind": "xlsx"})
    assert xl["file_format"] == "xlsx" and xl["path"] == xl["xlsx_path"]


async def test_report_projects_table_file_and_download(env, specs, prompts) -> None:
    meta = _meta()
    a, slug_a = await _seed(env, title="Первый", meta=meta, frames=3, images=2, videos=1, usd=0.5)
    _write_reply(slug_a, "excel_gpt_1", REPLY)
    b, _ = await _seed(env, title="Второй", meta={}, frames=1, status="new")
    ctx = FakeCtx(_routes(a, slug_a, meta, RUNS))
    run = specs.get("report.projects").executor

    out = await run(ctx, {"project_ids": [a, b, 999999], "kind": "json"})
    rows = {r["project_id"]: r for r in out["rows"]}
    ra, rb = rows[a], rows[b]
    assert ra["title"] == "Ролик про кота" and ra["frames"] == 3 and ra["images"] == 2 and ra["videos"] == 1
    assert ra["spent_usd"] == 0.5 and rb["spent_usd"] is None and rb["title"] == "Второй"
    assert ra["steps"].endswith(f"из {ra['steps_total']}") and ra["steps_done"] >= 1
    assert "Работа с GPT 1: Разметка кадров (razmetka) (готово)" in ra["gpt"]
    assert ra["gpt_nodes"][0]["frames_written"] == 2 and rb["gpt"] == "нет"
    assert rows[999999]["found"] is False and out["summary"]["missing"] == [999999]
    assert out["summary"]["frames"] == 4 and "Не найдены: #999999" in out["text_ru"]

    chat = await run(ctx, {"project_ids": [a, b]})
    assert chat["download_url"] == f"/api/projects/reports/{chat['file_name']}"
    assert chat["file_name"].startswith("projects_report_") and chat["file_format"] == "html"
    assert chat["path"].startswith(str(settings.data_dir / "reports"))
    assert "## Отчёт по проектам" in chat["markdown"] and "| № | Проект |" in chat["markdown"]
    xlsx_url = chat["download_url"][: -len(".html")] + ".xlsx"
    assert f"[Скачать отчёт (html)]({chat['download_url']})" in chat["text_ru"]
    assert f"[Скачать отчёт (xlsx)]({xlsx_url})" in chat["text_ru"]
    assert known_text("report.projects", chat) == chat["text_ru"]
    assert download_links(chat) == [f"[Скачать отчёт (html)]({chat['download_url']})",
                                    f"[Скачать отчёт (xlsx)]({xlsx_url})"]
    # Проект без записей трат: «нет данных», а не пустая ячейка и не 0.
    assert "| нет данных |" in chat["markdown"]
    wb = load_workbook(chat["xlsx_path"])
    ws = wb[tpl.PROJECTS_SHEET]
    assert [c.value for c in ws[1]] == [label for _k, label in tpl.PROJECTS_COLUMNS]
    assert ws.max_row == 3 and "Итог" in wb.sheetnames and tpl.GPT_SHEET in wb.sheetnames
    assert "нет данных" in [c.value for c in ws[3]]
    got = await env.client.get(xlsx_url)
    assert got.status_code == 200 and got.content[:2] == b"PK"
    page = await env.client.get(chat["download_url"])
    assert page.status_code == 200 and page.headers["content-security-policy"] == "sandbox"
    assert "Сводка по проектам" in page.text and f'id="project-{a}"' in page.text
    assert "Разметка кадров (razmetka)" in page.text and "Прогоны нод" in page.text
    (settings.data_dir / "reports" / "other.xlsx").write_bytes(b"PK")
    for bad in ("..%2Fstate.db", "other.xlsx", "projects_report_x.exe", "projects_report_none.xlsx"):
        got_bad = await env.client.get(f"/api/projects/reports/{bad}")
        # На этой линии раскодированный «../state.db» уходит мимо роутера отчётов в страницу-заглушку
        # студии (HTML, 200): базу он не отдаёт, это и проверяем.
        assert got_bad.status_code == 404 or (
            bad.startswith("..") and not got_bad.content.startswith(b"SQLite")
            and "text/html" in got_bad.headers.get("content-type", "")
        ), (bad, got_bad.status_code)

    every = await run(ctx, {"kind": "json"})
    assert {r["project_id"] for r in every["rows"]} == {a, b}


async def test_report_projects_defaults_to_last_twenty(env, specs) -> None:
    for i in range(22):
        await _seed(env, title=f"p{i}", meta={}, frames=0)
    out = await specs.get("report.projects").executor(FakeCtx({}), {"kind": "json"})
    assert len(out["rows"]) == tpl.MAX_PROJECTS and "последние 20 из 22" in out["text_ru"]


def test_report_projects_is_registered_read_only_and_planned() -> None:
    from app.orchestrator.action_contracts.contracts import CHAT_V1_ACTION_IDS
    from app.orchestrator.actions.executor import whitelist

    ensure_builtin_actions()
    spec = registry.get("report.projects")
    assert spec.risk == "read" and not spec.spends_money and not spec.destructive and not spec.needs_approval
    assert spec.title_ru == "Отчёт по нескольким проектам"
    assert spec.params_schema["properties"]["project_ids"]["maxItems"] == 20
    assert "report.projects" in CHAT_V1_ACTION_IDS and "report.projects" in whitelist()
    for text in ("сделай отчёт по проектам 5,6,7", "отчёт по всем проектам"):
        sel = select_catalog(text, None, registry, 12)
        assert "report.projects" in sel.ids, text
        assert sel.scores["report.projects"] >= sel.scores.get("report.nodes", 0), text
    assert "report.projects" in notes_for({"report.projects"})


def test_download_links_keep_strict_report_paths() -> None:
    assert download_links({"download_url": "/api/projects/5/reports/a.xlsx", "format": "xlsx"})
    assert download_links({"download_url": "/api/projects/reports/projects_report_1.xlsx"})
    assert download_links({"download_url": "/api/projects/reports/../x.xlsx"}) == []
    assert download_links({"download_url": "/api/projects/5/reports/a b.xlsx"}) == []
    assert download_links({"download_url": "https://evil/x"}) == []


def test_template_is_the_single_place_for_columns() -> None:
    keys = {k for k, _l in tpl.GPT_NODE_COLUMNS}
    row_keys = {"label", "prompt_title", "model", "status_ru", "frames_written", "fields_ru",
                "characters_count", "scenes_count", "cost_usd", "cost_how", "source", "note"}
    assert keys <= row_keys
    assert tpl.field_label("image_prompt") == "промт картинки" and tpl.field_label("x_custom") == "x_custom"
    assert json.dumps(tpl.PROJECTS_COLUMNS, ensure_ascii=False)


async def test_report_projects_names_the_project_and_rounds_money(env, specs, prompts) -> None:
    meta = _meta()
    a, slug_a = await _seed(env, meta=meta, frames=4, usd=0.4415)
    ctx = FakeCtx(_routes(a, slug_a, meta, RUNS))
    out = await specs.get("report.projects").executor(ctx, {"project_ids": [a], "kind": "json"})
    assert f"#{a} «Ролик про кота»" in out["text_ru"]
    assert "траты $0.44 (оценка)" in out["text_ru"] and "0.4415" not in out["text_ru"]
