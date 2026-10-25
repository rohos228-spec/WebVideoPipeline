"""Отчёты по данным нод «Работа с GPT» и сводный отчёт по нескольким проектам.

Цифры собирает код (без модели). Вид таблиц и подписи — в report_template.py, HTML — в report_html.py.
Что нода записала, берём из всех пачек её последнего прогона (node_trace/<нода>/<прогон>/call_*):
ответ GPT и итоговые ops после правок кода (ops_final.json). Нет дневника прогона — из сохранённого
ответа excel_gpt_uploads/<нода>/gpt_reply.txt. Разбор тот же, что применяет ответ к базе, но без записи.
Архивные кадры (attrs.archived) в отчёт не попадают.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy.exc import SQLAlchemyError

from app.services.action_layer import report_template as tpl
from app.services.action_layer.context import ActionCallError, ActionContext

_REPLY_NAME = "gpt_reply.txt"
_REPLY_MAX_BYTES = 4 * 1024 * 1024
_NAMED_IN_HEAD = 3
_CHAT_LINE_MAX = 220
# Больше в одну ячейку xlsx Excel не пускает (32767); длинное значение идёт продолжением в следующих строках.
XLSX_CELL_MAX = 32000


# ── Общие помощники рендера ──


def clip(value: Any, n: int = 200) -> str:
    text = " ".join(str(value if value is not None else "").split())
    return text if len(text) <= n else text[: n - 1] + "…"


def md_cell(value: Any) -> str:
    return clip(value, 160).replace("|", "/") if value not in (None, "") else "—"


def markdown_table(rows: list[dict], columns: tuple[tuple[str, str], ...]) -> str:
    head = "| " + " | ".join(label for _k, label in columns) + " |"
    sep = "|" + "|".join("---" for _ in columns) + "|"
    body = ["| " + " | ".join(md_cell(r.get(k)) for k, _l in columns) + " |" for r in rows]
    return "\n".join([head, sep, *body])


def xlsx_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "@")):
        return " " + value
    if isinstance(value, (list, tuple, set)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    return value


def _xlsx_lines(values: list[Any]) -> list[list[Any]]:
    """Строка таблицы → строки листа: длинный текст режется на куски, ничего не теряется."""
    parts = [
        [v[i : i + XLSX_CELL_MAX] for i in range(0, len(v), XLSX_CELL_MAX)]
        if isinstance(v, str) and len(v) > XLSX_CELL_MAX
        else [v]
        for v in values
    ]
    extra = max(len(p) for p in parts) if parts else 1
    lines = [[p[0] for p in parts]]
    for i in range(1, extra):
        line = [p[i] if i < len(p) else None for p in parts]
        if line[0] is None:
            line[0] = f"{values[0]} (продолжение)"
        lines.append(line)
    return lines


def write_workbook(
    path: Path,
    sheets: list[tuple[str, tuple[tuple[str, str], ...], list[dict]]],
    summary: list[tuple[str, Any]],
    *,
    widths: dict[str, dict[str, int]] | None = None,
) -> None:
    """Один xlsx: лист «Итог» последним, на каждом листе шапка жирным и закреплена."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    first = True
    for name, columns, rows in sheets:
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        ws.append([label for _k, label in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            for line in _xlsx_lines([xlsx_value(r.get(k)) for k, _l in columns]):
                ws.append(line)
        fixed = (widths or {}).get(name) or {}
        for i, (key, label) in enumerate(columns, start=1):
            letter = get_column_letter(i)
            longest = max([len(str(label))] + [len(str(xlsx_value(r.get(key)) or "")) for r in rows[:200]])
            ws.column_dimensions[letter].width = fixed.get(letter) or min(
                max(10, longest + 2), tpl.COLUMN_WIDTH_MAX
            )
        ws.freeze_panes = "A2"
    total = wb.active if first else wb.create_sheet(tpl.SUMMARY_SHEET)
    total.title = tpl.SUMMARY_SHEET
    for name, value in summary:
        total.append([name, xlsx_value(value)])
    total.column_dimensions["A"].width = 36
    total.column_dimensions["B"].width = 60
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _stamp() -> str:
    """Метка имени файла отчёта: время и случайный хвост.

    Только секунды не уникальны: два вызова в одну секунду (два клиента студии в общей
    папке отчётов) писали бы в один файл и перетирали друг друга.
    """
    import uuid

    return f"{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}"


# ── Что записала нода «Работа с GPT» ──


def read_reply(data_dir: str | Path, node_key: str) -> str | None:
    """Сохранённый ответ ноды; только внутри папки проекта, без обхода путей."""
    root = Path(data_dir).resolve()
    base = (root / "excel_gpt_uploads").resolve()
    path = (base / str(node_key) / _REPLY_NAME).resolve()
    if path.parent.parent != base or not path.is_file():
        return None
    try:
        if path.stat().st_size > _REPLY_MAX_BYTES:
            return None
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return None


def _read_text(path: Path) -> str | None:
    try:
        if not path.is_file() or path.stat().st_size > _REPLY_MAX_BYTES:
            return None
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return None


def _read_json(path: Path) -> Any:
    import json

    text = _read_text(path)
    if text is None:
        return None
    try:
        return json.loads(text)
    except ValueError:
        return None


def _call_order(path: Path) -> tuple[int, str]:
    parts = path.name.split("_")
    return (int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else 0, path.name)


def read_batches(data_dir: str | Path, node_key: str) -> dict:
    """Все пачки последнего прогона ноды с вызовами GPT; только внутри node_trace проекта."""
    from app.services.node_trace import _safe_name, read_run, trace_root

    root = trace_root(data_dir).resolve()
    node_dir = (root / _safe_name(str(node_key))).resolve()
    empty: dict[str, Any] = {"run": {}, "calls": []}
    if node_dir.parent != root or not node_dir.is_dir():
        return empty
    try:
        runs = sorted(
            p
            for p in node_dir.iterdir()
            if p.is_dir() and not p.name.endswith(".del") and any(p.glob("call_*"))
        )
    except OSError:
        return empty
    if not runs:
        return empty
    run_dir = runs[-1]
    calls = []
    for call_dir in sorted((p for p in run_dir.glob("call_*") if p.is_dir()), key=_call_order):
        ops_final = _read_json(call_dir / "ops_final.json")
        calls.append(
            {
                "name": call_dir.name,
                "reply": _read_text(call_dir / "reply.txt"),
                "ops_final": ops_final if isinstance(ops_final, list) else None,
                "error": (_read_text(call_dir / "error.txt") or "").strip(),
            }
        )
    return {"run": read_run(run_dir), "calls": calls}


def _canon_field(raw: Any) -> str:
    from app.services.db_apply import FIELD_ALIASES

    key = str(raw).strip()
    return FIELD_ALIASES.get(key.lower().replace(" ", "_")) or key


@dataclass
class FrameIndex:
    """Кадры проекта для отчёта: живые (uuid → номер) и архивные (их в отчёт не берём)."""

    num_by_uuid: dict[str, int] = field(default_factory=dict)
    archived_uuids: set[str] = field(default_factory=set)
    archived_numbers: set[int] = field(default_factory=set)
    frames: list[dict] = field(default_factory=list)

    def is_archived(self, ref: Any) -> bool:
        text = str(ref if ref is not None else "").strip()
        if text in self.archived_uuids:
            return True
        if text.isdigit():
            n = int(text)
            return n >= ARCHIVE_NUMBER_FROM or (
                n in self.archived_numbers and n not in self.num_by_uuid.values()
            )
        return False


# Парковочный номер архивного кадра (db_v2.frame_is_archived).
ARCHIVE_NUMBER_FROM = 1_000_000


def _index(index: FrameIndex | dict[str, int] | None) -> FrameIndex:
    if isinstance(index, FrameIndex):
        return index
    return FrameIndex(num_by_uuid=dict(index or {}))


def _frame_number(ref: Any, num_by_uuid: dict[str, int]) -> int | str:
    text = str(ref if ref is not None else "").strip()
    if text in num_by_uuid:
        return num_by_uuid[text]
    if text.isdigit():
        return int(text)
    return f"? {text[:12]}" if text else "?"


def _card_name(card: Any, *keys: str) -> str:
    if not isinstance(card, dict):
        return clip(card, 60)
    parts = [clip(card.get(k), 60) for k in keys if str(card.get(k) or "").strip()]
    return " — ".join(dict.fromkeys(parts)) or "без названия"


def cell_text(value: Any) -> str:
    """Значение поля целиком: текст как есть, список и словарь — JSON без сокращений."""
    import json

    if value is None:
        return ""
    if isinstance(value, str):
        return value
    try:
        return json.dumps(value, ensure_ascii=False)
    except (TypeError, ValueError):
        return str(value)


def _empty_writes() -> dict[str, Any]:
    return {
        "parsed": False,
        "frames": {},
        "values": {},
        "project_fields": [],
        "replace_frames": None,
        "characters": [],
        "scenes": [],
        "partial": False,
        "archived_skipped": 0,
    }


def _apply_data(out: dict[str, Any], data: dict, index: FrameIndex) -> None:
    """Одна пачка ответа поверх уже собранного: поздняя пачка перезаписывает то же поле кадра."""
    out["parsed"] = True
    out["partial"] = out["partial"] or bool(data.get("_salvaged_partial"))
    ops = data.get("ops") if isinstance(data.get("ops"), list) else data.get("actions")
    for op in ops if isinstance(ops, list) else []:
        if not isinstance(op, dict):
            continue
        target = op.get("target") or "frame"
        fields = op.get("fields") if isinstance(op.get("fields"), dict) else {}
        names = [_canon_field(k) for k in fields]
        if target == "frame":
            ref = op.get("frame_uuid") or op.get("frame_number") or op.get("кадр")
            if index.is_archived(ref):
                out["archived_skipped"] += 1
                continue
            num = _frame_number(ref, index.num_by_uuid)
            known = out["frames"].setdefault(num, [])
            known.extend(n for n in names if n not in known)
            slot = out["values"].setdefault(num, {})
            for raw, value in fields.items():
                slot[_canon_field(raw)] = cell_text(value)
        elif target == "project":
            out["project_fields"].extend(n for n in names if n not in out["project_fields"])
        elif target == "replace_frames":
            specs = op.get("frames") or op.get("кадры") or []
            out["replace_frames"] = (out["replace_frames"] or 0) + (
                len(specs) if isinstance(specs, list) else 0
            )
    chars = data.get("characters") if isinstance(data.get("characters"), list) else []
    for c in chars:
        name = _card_name(c, "id", "code", "имя", "name")
        if name not in out["characters"]:
            out["characters"].append(name)
    scenes = data.get("scenes") if isinstance(data.get("scenes"), list) else []
    for sc in scenes:
        name = _card_name(sc, "id_scene", "id", "scene_id", "место", "place", "title")
        if name not in out["scenes"]:
            out["scenes"].append(name)


def _parse(text: str | None) -> dict | None:
    if not text:
        return None
    from app.services.db_apply import extract_apply_ops_json

    try:
        data = extract_apply_ops_json(text)
    except Exception:  # noqa: BLE001 - кривой ответ не роняет отчёт
        data = None
    return data if isinstance(data, dict) else None


def _sorted_frames(out: dict[str, Any]) -> dict[str, Any]:
    out["frames"] = dict(sorted(out["frames"].items(), key=lambda kv: (isinstance(kv[0], str), kv[0])))
    return out


def writes_from_reply(text: str | None, num_by_uuid: FrameIndex | dict[str, int] | None) -> dict:
    """Разбор одного ответа: кадр → поля, поля проекта, персонажи, сцены. Ничего не пишет."""
    out = _empty_writes()
    data = _parse(text)
    if data is not None:
        _apply_data(out, data, _index(num_by_uuid))
    return _sorted_frames(out)


def writes_from_batches(batches: dict, index: FrameIndex | dict[str, int] | None) -> dict:
    """Все пачки прогона по порядку. Итоговые ops (после правок кода) важнее сырого ответа GPT."""
    out = _empty_writes()
    idx = _index(index)
    calls = batches.get("calls") or []
    out.update(batches=len(calls), batches_unparsed=0, batch_errors=0)
    for call in calls:
        data = _parse(call.get("reply"))
        if isinstance(call.get("ops_final"), list):
            data = {**(data or {}), "ops": call["ops_final"]}
        if call.get("error"):
            out["batch_errors"] += 1
        if data is None:
            out["batches_unparsed"] += 1
            continue
        _apply_data(out, data, idx)
    return _sorted_frames(out)


def _labels_ru(names: list[str]) -> str:
    return ", ".join(tpl.field_label(n) for n in names)


def prompt_title(variant: str | None) -> str:
    if not variant:
        return ""
    from app.services import prompt_library as pl

    try:
        titles = pl.load_file_meta(pl.EXCEL_GPT_UNIFIED_STEP)
    except Exception:  # noqa: BLE001
        titles = {}
    info = titles.get(variant) if isinstance(titles.get(variant), dict) else {}
    title = str(info.get("title") or "").strip()
    if not title or title == variant or variant in title:
        return title or variant
    return f"{title} ({variant})"


def _frame_row(fr: Any) -> dict:
    attrs = fr.attrs if isinstance(fr.attrs, dict) else {}
    return {
        "number": int(fr.number),
        "uuid": str(fr.uuid or ""),
        "voiceover_text": fr.voiceover_text or "",
        "place": str(attrs.get("место") or attrs.get("place") or ""),
        "attrs": attrs,
        "image_prompt": str(getattr(fr, "image_prompt", None) or ""),
        "animation_prompt": str(getattr(fr, "animation_prompt", None) or ""),
    }


async def frame_index(pid: int, data_dir: str | Path | None = None) -> FrameIndex:
    """Кадры проекта из единой базы (``data_dir`` — для совместимости вызова, project.db на этой линии нет)."""
    from sqlalchemy import select

    from app.models import Frame
    from app.services.action_layer import builtin_insight as bi
    from app.services.db_v2 import frame_is_archived

    try:
        async with bi._sessions()() as s:
            rows = (await s.scalars(select(Frame).where(Frame.project_id == pid))).all()
    except SQLAlchemyError:
        rows = []
    out = FrameIndex()
    for fr in sorted(rows, key=lambda f: int(f.number or 0)):
        if frame_is_archived(fr):
            if fr.uuid:
                out.archived_uuids.add(str(fr.uuid))
            out.archived_numbers.add(int(fr.number or 0))
            continue
        if fr.uuid:
            out.num_by_uuid[str(fr.uuid)] = int(fr.number)
        out.frames.append(_frame_row(fr))
    return out


def is_gpt_node(node: dict | None, node_type: str = "") -> bool:
    from app.services.excel_gpt_node import is_excel_gpt_node_type

    raw = str((node or {}).get("type") or node_type or "")
    return is_excel_gpt_node_type(raw)


def node_writes(data_dir: str | Path, node_key: str, index: FrameIndex | dict[str, int] | None) -> dict:
    """Что записала нода: все пачки прогона из node_trace, без него — сохранённый ответ."""
    batches = read_batches(data_dir, node_key)
    if batches["calls"]:
        writes = writes_from_batches(batches, index)
        writes["run"] = batches["run"]
        writes["source"] = f"пачек прогона: {len(batches['calls'])}"
        return writes
    writes = writes_from_reply(read_reply(data_dir, node_key), index)
    writes.update(run={}, source="последний ответ" if writes["parsed"] else "")
    return writes


def _writes_note(writes: dict, status_ru: str) -> str:
    notes: list[str] = []
    if not writes["parsed"]:
        if status_ru != "не запускалась":
            notes.append("сохранённого ответа с данными нет")
    elif writes["partial"]:
        notes.append("ответ был обрезан, разобрана целая часть")
    if writes.get("batches_unparsed"):
        notes.append(f"пачек без разбора: {writes['batches_unparsed']}")
    if writes.get("batch_errors"):
        notes.append(f"пачек с ошибкой: {writes['batch_errors']}")
    skipped = (writes.get("run") or {}).get("skipped") or []
    if skipped:
        notes.append(f"пачек не записано: {len(skipped)}")
    if writes.get("archived_skipped"):
        notes.append(f"записей в архивные кадры (не показаны): {writes['archived_skipped']}")
    if writes["project_fields"]:
        notes.append("проект: " + _labels_ru(writes["project_fields"]))
    if writes["replace_frames"] is not None:
        notes.append(f"пересобрала кадры: {writes['replace_frames']}")
    return "; ".join(notes)


def gpt_nodes(
    project: Any, detail: dict, cards: dict[str, dict], index: FrameIndex | dict[str, int] | None
) -> list[dict]:
    """Строки по нодам «Работа с GPT» в порядке канваса."""
    from app.services.action_layer import builtin_insight as bi
    from app.services.action_layer.builtin_group import node_label_ru

    meta = bi._meta(detail)
    out: list[dict] = []
    for node in bi._canvas_nodes(meta):
        key = str(node["id"])
        if not is_gpt_node(node):
            continue
        card = cards.get(key) or {}
        variant = (card.get("prompt") or {}).get("variant") if card else None
        if not card:
            raw, eff = bi._node_types(node)
            variant = (bi._prompt_for(detail, key, raw, eff) or {}).get("variant")
        status_ru = card.get("status_ru") or "не запускалась"
        writes = node_writes(project.data_dir, key, index)
        fields: list[str] = []
        for names in writes["frames"].values():
            fields.extend(n for n in names if n not in fields)
        run = writes.get("run") or {}
        out.append(
            {
                "node_key": key,
                "label": node_label_ru(meta, key),
                "prompt_variant": variant or "",
                "prompt_title": prompt_title(variant),
                "model": card.get("model") or "",
                "status_ru": status_ru,
                "frames_written": len(writes["frames"]),
                "fields": fields,
                "fields_ru": _labels_ru(fields),
                "frames": {str(k): v for k, v in writes["frames"].items()},
                "values": {str(k): v for k, v in (writes.get("values") or {}).items()},
                "characters": writes["characters"],
                "characters_count": len(writes["characters"]),
                "scenes": writes["scenes"],
                "scenes_count": len(writes["scenes"]),
                "project_fields": writes["project_fields"],
                "note": _writes_note(writes, status_ru),
                "source": writes.get("source") or "",
                "batches": int(writes.get("batches") or 0),
                "archived_skipped": int(writes.get("archived_skipped") or 0),
                "run_when": run.get("finished") or run.get("started") or "",
                "run_status": run.get("status") or "",
                "run_calls": run.get("calls"),
                "run_counts": run.get("counts") or {},
                "run_dir": run.get("dir") or "",
                "cost_usd": card.get("cost_usd"),
                "cost_how": card.get("cost_how") or "",
            }
        )
    return out


def gpt_frame_rows(nodes: list[dict]) -> list[dict]:
    """Поле кадра — строка, значение целиком. Поле, которое потом записала другая нода, помечено:
    в кадре осталось её значение, а это пропало («последний побеждает»)."""
    later: dict[tuple[str, str], str] = {}
    rows: list[dict] = []
    for n in reversed(nodes):
        for frame, names in n["frames"].items():
            vals = (n.get("values") or {}).get(str(frame)) or {}
            for name in names:
                key = (str(frame), name)
                rows.append(
                    {
                        "label": n["label"],
                        "frame": frame,
                        "field": name,
                        "field_ru": tpl.field_label(name),
                        "value": vals.get(name, ""),
                        "overwritten": f"да, потом «{later[key]}»" if key in later else "",
                    }
                )
            for name in names:
                later.setdefault((str(frame), name), n["label"])
    order = {n["label"]: i for i, n in enumerate(nodes)}
    return sorted(
        rows, key=lambda r: (order.get(r["label"], 0), (isinstance(r["frame"], str), r["frame"]), r["field"])
    )


def overwrite_text(nodes: list[dict]) -> str:
    """По-русски, кто чьи поля перезаписал: нода в порядке канваса, поле кадра, которое до неё
    записала другая нода «Работа с GPT». Тот же расчёт, что колонка «Перезаписано» в html."""
    if not nodes:
        return ""
    # Ноды сравниваем по ключу, а не по названию: две ноды «Работа с GPT» с одним названием
    # тоже перезаписывают друг друга, как в колонке html.
    last: dict[tuple[str, str], tuple[str, str]] = {}
    lines = ["Кто чьи поля перезаписал (работа с GPT по порядку на канвасе):"]
    for n in nodes:
        me = str(n.get("node_key") or n["label"])
        frames = n.get("frames") or {}
        if not frames:
            lines.append(f"- {n['label']}: пока ничего не записала в кадры, поэтому ничего не перезаписала.")
            continue
        over: dict[str, dict[str, set[str]]] = {}
        for frame, names in frames.items():
            for name in names:
                prev = last.get((str(frame), name))
                if prev and prev[0] != me:
                    who = f"«{prev[1]}»" + (
                        " (другая нода с тем же названием)" if prev[1] == n["label"] else ""
                    )
                    over.setdefault(who, {}).setdefault(name, set()).add(str(frame))
        if not over:
            lines.append(
                f"- {n['label']}: записала {_labels_ru(list(n.get('fields') or []))} "
                f"в кадрах {len(frames)}, чужих полей не перезаписала."
            )
        else:
            parts = []
            for prev, fields in over.items():
                named = ", ".join(
                    f"{tpl.field_label(name)} (кадров {len(fr)})" for name, fr in fields.items()
                )
                parts.append(f"записанные {prev}: {named}")
            lines.append(f"- {n['label']} перезаписала поля, " + "; ".join(parts) + ".")
        for frame, names in frames.items():
            for name in names:
                last[(str(frame), name)] = (me, n["label"])
    return "\n".join(lines)


def gpt_list_rows(nodes: list[dict]) -> list[dict]:
    rows = [{"label": n["label"], "kind": "персонаж", "name": c} for n in nodes for c in n["characters"]]
    rows += [{"label": n["label"], "kind": "сцена", "name": s} for n in nodes for s in n["scenes"]]
    return rows


def gpt_text(nodes: list[dict]) -> str:
    """Короткий текст сверху: по одной строке на ноду."""
    if not nodes:
        return ""
    done = sum(1 for n in nodes if n["status_ru"] == "готово")
    lines = [f"{tpl.GPT_SECTION_TITLE}: нод {len(nodes)}, готово {done}."]
    for n in nodes:
        parts = [f"промт «{n['prompt_title']}»" if n["prompt_title"] else "", n["model"], n["status_ru"]]
        what = []
        if n["frames_written"]:
            what.append(
                f"заполнила кадров {n['frames_written']}"
                + (f" ({clip(n['fields_ru'], 80)})" if n["fields_ru"] else "")
            )
        if n["characters_count"]:
            what.append(f"персонажей {n['characters_count']}")
        if n["scenes_count"]:
            what.append(f"сцен {n['scenes_count']}")
        line = f"{n['label']}: " + ", ".join(p for p in parts if p)
        line += (". " + ", ".join(what) + ".") if what else "."
        lines.append(clip(line, _CHAT_LINE_MAX))
    return "\n".join(lines)


async def project_counts(pid: int, meta: dict) -> dict:
    """Персонажи в базе и сцены в реестре — для итога отчёта."""
    from sqlalchemy import func, select

    from app.models import Entity
    from app.services.action_layer import builtin_insight as bi

    chars = None
    try:
        async with bi._sessions()() as s:
            chars = int(
                await s.scalar(
                    select(func.count(Entity.id)).where(Entity.project_id == pid, Entity.type == "character")
                )
                or 0
            )
    except SQLAlchemyError:
        chars = None
    reg = meta.get("scene_registry")
    scenes = (
        sum(1 for r in reg if not (isinstance(r, dict) and r.get("активен") is False))
        if isinstance(reg, list)
        else 0
    )
    return {"characters_in_project": chars, "scenes_in_registry": scenes}


# ── Отчёт по нескольким проектам ──


def shared_reports_dir() -> Path:
    from app.settings import settings

    return Path(settings.data_dir) / "reports"


async def _project_ids(params: dict) -> tuple[list[int], int]:
    from sqlalchemy import func, select

    from app.models import Project
    from app.services.action_layer import builtin_insight as bi

    ids = [int(x) for x in params.get("project_ids") or []]
    if ids:
        return list(dict.fromkeys(ids))[: tpl.MAX_PROJECTS], len(set(ids))
    async with bi._sessions()() as s:
        total = int(await s.scalar(select(func.count(Project.id))) or 0)
        rows = (await s.scalars(select(Project.id).order_by(Project.id.desc()).limit(tpl.MAX_PROJECTS))).all()
    return sorted(int(r) for r in rows), total


async def _grouped(pids: list[int]) -> dict[str, dict[int, float]]:
    from types import SimpleNamespace

    from sqlalchemy import func, select, text

    from app.models import Artifact, ArtifactKind, Frame
    from app.services.action_layer import builtin_insight as bi
    from app.services.db_v2 import frame_is_archived

    out: dict[str, dict[int, float]] = {"frames": {}, "images": {}, "videos": {}, "spent": {}}
    if not pids:
        return out
    from app.services.frame_store import session_factory_for

    async def count(s, ids: list[int]) -> None:
        # Счёт по проектам из единой базы; max — на случай повторного прохода по тем же id.
        # Архивные кадры (attrs.archived, парковочный номер) не считаем.
        live: dict[int, int] = {}
        for pid, number, attrs in (
            await s.execute(
                select(Frame.project_id, Frame.number, Frame.attrs).where(Frame.project_id.in_(ids))
            )
        ).all():
            if not frame_is_archived(SimpleNamespace(number=number, attrs=attrs)):
                live[int(pid)] = live.get(int(pid), 0) + 1
        for pid, n in live.items():
            out["frames"][pid] = max(out["frames"].get(pid, 0), n)
        for pid, kind, n in (
            await s.execute(
                select(Artifact.project_id, Artifact.kind, func.count(Artifact.id))
                .where(
                    Artifact.project_id.in_(ids),
                    Artifact.kind.in_([ArtifactKind.scene_image, ArtifactKind.scene_video]),
                )
                .group_by(Artifact.project_id, Artifact.kind)
            )
        ).all():
            kind_v = kind.value if hasattr(kind, "value") else str(kind)
            bucket = "images" if kind_v == ArtifactKind.scene_image.value else "videos"
            out[bucket][int(pid)] = max(out[bucket].get(int(pid), 0), int(n))

    for pid in pids:
        try:
            async with (await session_factory_for(pid))() as s:
                await count(s, [pid])
        except Exception:  # noqa: BLE001
            continue
    from app.services.action_layer.plan import member_view

    # Траты — себестоимость из общей кассы провайдеров: участнику не видны.
    hide_spent = member_view()
    try:
        async with bi._sessions()() as s:
            await count(s, pids)
            if hide_spent:
                return out
            marks = ", ".join(f":p{i}" for i in range(len(pids)))
            rows = (
                await s.execute(
                    text(
                        f"SELECT project_id, SUM(usd) FROM spend_ledger WHERE project_id IN ({marks}) "
                        "GROUP BY project_id"
                    ),
                    {f"p{i}": p for i, p in enumerate(pids)},
                )
            ).all()
            for pid, usd in rows:
                out["spent"][int(pid)] = float(usd or 0.0)
    except SQLAlchemyError:
        pass
    return out


async def _project_detail(ctx: ActionContext, pid: int, row: Any) -> dict:
    from app.services.action_layer.builtin_group import _effective_project

    try:
        detail = await _effective_project(ctx, pid)
        if isinstance(detail, dict):
            return detail
    except ActionCallError:
        pass
    return {
        "id": pid,
        "title": row.title,
        "topic": row.topic,
        "meta": dict(row.meta or {}),
        "prompt_overrides": dict(row.prompt_overrides or {}),
    }


_PROJECT_GPT_KEYS = (
    "node_key",
    "label",
    "prompt_variant",
    "prompt_title",
    "model",
    "status_ru",
    "frames_written",
    "characters_count",
    "scenes_count",
    "cost_usd",
    "cost_how",
    "note",
)


async def _project_row_report(ctx: ActionContext, pid: int, row: Any, counts: dict) -> dict:
    from app.services.action_layer import builtin_insight as bi
    from app.services.action_layer import report_costs
    from app.services.action_layer.builtin_group import node_label_ru

    if row is None:
        return {
            "project_id": pid,
            "title": "проект не найден",
            "status_ru": "—",
            "found": False,
            "gpt_nodes": [],
        }
    detail = await _project_detail(ctx, pid, row)
    meta = bi._meta(detail)
    status = bi._row_status(row) or str(detail.get("status") or "")
    steps = bi._pipeline(status, meta)
    done = sum(1 for s in steps if s["state"] == "готово")
    runs = {str(nr.get("node_key")): nr for nr in await bi._node_runs(ctx, pid)}
    by_key = {str(n["id"]): n for n in bi._canvas_nodes(meta)}
    project = bi._transient(pid, detail, row)
    split = await report_costs.node_costs(
        pid, node_runs=list(runs.values()), step_nodes=report_costs.step_nodes_of(list(by_key.values()))
    )
    cards: dict[str, dict] = {}
    for key, node in by_key.items():
        if is_gpt_node(node):
            cards[key] = bi._node_card(
                project,
                detail,
                key,
                node,
                runs.get(key),
                checks=[],
                cost=split.by_node.get(key),
                detailed=False,
            )
            cards[key]["cost_how"] = ", ".join(split.how.get(key) or [])
    nodes = gpt_nodes(project, detail, cards, await frame_index(pid, project.data_dir))
    gpt = "; ".join(
        f"{n['label']}: {n['prompt_title'] or 'промт по умолчанию'} ({n['status_ru']})" for n in nodes
    )
    trace_runs, _diary = bi._trace_runs(project.data_dir)
    spent = counts["spent"].get(pid)
    return {
        "project_id": pid,
        "found": True,
        "title": clip(detail.get("title") or row.title or row.topic, 80),
        "status": status,
        "status_ru": bi._status_ru(status),
        "steps": f"{done} из {len(steps)}",
        "steps_done": done,
        "steps_total": len(steps),
        "gpt": gpt or "нет",
        "gpt_nodes": [{k: n.get(k) for k in _PROJECT_GPT_KEYS} for n in nodes],
        "frames": int(counts["frames"].get(pid, 0)),
        "images": int(counts["images"].get(pid, 0)),
        "videos": int(counts["videos"].get(pid, 0)),
        "spent_usd": None if spent is None else round(spent, 4),
        "spent_by_source": {k: round(v, 4) for k, v in split.by_source.items()},
        "attributed_usd": round(split.attributed, 4),
        "runs": trace_runs,
        "labels": {key: node_label_ru(meta, key) for key in by_key},
        "updated_at": bi._iso(getattr(row, "updated_at", None)),
    }


_SPENT_ROW_KEYS = ("spent_usd", "spent_by_source", "attributed_usd")
_SPENT_NODE_KEYS = ("cost_usd", "cost_how")


def _member_row(row: dict) -> dict:
    """Строка проекта для участника: без сумм в $ (себестоимость из общей кассы)."""
    out = {k: v for k, v in row.items() if k not in _SPENT_ROW_KEYS}
    if isinstance(out.get("gpt_nodes"), list):
        out["gpt_nodes"] = [
            {k: v for k, v in n.items() if k not in _SPENT_NODE_KEYS} if isinstance(n, dict) else n
            for n in out["gpt_nodes"]
        ]
    return out


MEMBER_NO_FILE_TEXT = (
    "Файл отчёта по проектам здесь не создаётся: общие файлы отчётов студии скачивает админ. "
    "Вся таблица — ниже в чате."
)


def _without_spent(columns: tuple[tuple[str, str], ...]) -> tuple[tuple[str, str], ...]:
    return tuple(col for col in columns if col[0] not in ("spent_usd", "cost_usd"))


def _project_line(r: dict, *, show_spent: bool = True) -> str:
    """Строка проекта для чата. Участнику — без трат (суммы в $ видит только админ)."""
    from app.services.action_layer import builtin_insight as bi

    pid = r["project_id"]
    if not r.get("found"):
        return f"- #{pid}: проект не найден."
    spent = r.get("spent_usd")
    money = f"траты {bi._usd_text(spent)} (оценка)" if spent is not None else "траты: нет данных"
    return (
        f"- #{pid} «{r.get('title')}»: {r.get('status_ru') or 'статус неизвестен'}, шагов готово "
        f"{r.get('steps')}, кадров {r.get('frames')}, картинок {r.get('images')}, видео {r.get('videos')}"
        + (f", {money}." if show_spent else ".")
    )


def with_no_data(rows: list[dict], keys: tuple[str, ...]) -> list[dict]:
    """Для файла и чата: пустая стоимость — «нет данных», не пустая ячейка и не 0.
    Нода, которая не запускалась, ничего не тратила: у неё прочерк."""
    from app.services.action_layer.report_costs import empty_cost

    return [{**r, **{k: empty_cost(r.get("status_ru")) for k in keys if r.get(k) is None}} for r in rows]


async def report_projects(ctx: ActionContext, params: dict) -> dict:
    from app.models import Project
    from app.services.action_layer import builtin_insight as bi
    from app.services.action_layer import report_html
    from app.services.action_layer.plan import member_view

    fmt = params.get("kind", "chat")
    show_spent = not member_view()
    pids, total = await _project_ids(params)
    counts = await _grouped(pids)
    rows: list[dict] = []
    for pid in pids:
        try:
            async with bi._sessions()() as s:
                row = await s.get(Project, pid)
        except SQLAlchemyError:
            row = None
        rows.append(await _project_row_report(ctx, pid, row, counts))
    if not show_spent:
        # Участнику — без трат у самого источника: строки уходят и в kind="json"
        # (ответ шага, журнал действий), а не только в отрисованную таблицу.
        rows = [_member_row(r) for r in rows]
    found = [r for r in rows if r.get("found")]
    spent = sum(r.get("spent_usd") or 0.0 for r in found)
    missing = [r["project_id"] for r in rows if not r.get("found")]
    # Номера и названия в первой строке: без них ответ на «отчёт по проекту 6» говорил,
    # что не может подтвердить, что это проект #6.
    named = ", ".join(f"#{r['project_id']} «{r['title']}»" for r in found[:_NAMED_IN_HEAD])
    head = (
        f"{tpl.PROJECTS_TITLE}: {len(found)}"
        + (
            f" (последние {len(pids)} из {total})"
            if not params.get("project_ids") and total > len(pids)
            else ""
        )
        + (f" — {named}" if named and len(found) <= _NAMED_IN_HEAD else "")
        + f". Кадров {sum(r['frames'] for r in found)}, картинок {sum(r['images'] for r in found)}, "
        f"видео {sum(r['videos'] for r in found)}"
        + (f", траты {bi._usd_text(spent)} (оценка)." if show_spent else ".")
    )
    if missing:
        head += " Не найдены: " + ", ".join(f"#{p}" for p in missing) + "."
    summary = {
        "projects": len(found),
        "missing": missing,
        "frames": sum(r["frames"] for r in found),
        "images": sum(r["images"] for r in found),
        "videos": sum(r["videos"] for r in found),
        "spent_usd_estimate": round(spent, 4) if show_spent else None,
    }
    if not show_spent:
        summary.pop("spent_usd_estimate")
    out: dict[str, Any] = {"format": fmt, "project_ids": pids, "summary": summary}
    if fmt == "json":
        out.update(rows=rows, text_ru=head)
        return bi._safe(out)
    pairs = [
        ("Отчёт", tpl.PROJECTS_TITLE),
        ("Сформирован", datetime.now().isoformat(timespec="seconds")),
        ("Проектов", len(found)),
        ("Кадров (без архивных)", summary["frames"]),
        ("Картинок", summary["images"]),
        ("Видео", summary["videos"]),
    ]
    if show_spent:
        pairs.append(("Траты, $ (оценка)", summary["spent_usd_estimate"]))
    columns = tpl.PROJECTS_COLUMNS if show_spent else _without_spent(tpl.PROJECTS_COLUMNS)
    chat_columns = tpl.PROJECTS_CHAT_COLUMNS if show_spent else _without_spent(tpl.PROJECTS_CHAT_COLUMNS)
    if missing:
        pairs.append(("Не найдены", ", ".join(f"#{p}" for p in missing)))
    if not show_spent:
        # Участник: файлы общих отчётов студии (/api/projects/reports/...) отдаются только админу
        # (project_reports.download_shared_report), ссылка вела бы на 403. Файл не пишем, вся
        # таблица — в чате.
        table = markdown_table(with_no_data(rows, ("spent_usd",)), chat_columns)
        gpt_lines = [
            f"- #{r['project_id']}: {clip(r['gpt'], 160)}"
            for r in rows
            if r.get("gpt") not in (None, "", "нет")
        ]
        gpt_text_md = ("\n\n" + tpl.GPT_SECTION_TITLE + ":\n" + "\n".join(gpt_lines)) if gpt_lines else ""
        out.update(rows=len(rows), file_format=None)
        out["markdown"] = f"## {tpl.PROJECTS_TITLE}\n\n{head}\n\n{table}{gpt_text_md}"
        # Строка на каждый проект отчёта, как у админа: при 10 строках из 14 ответ говорил
        # «подробности о 10 из 14». Файла у участника нет — остаток назван, вся таблица выше.
        per = [_project_line(r, show_spent=False) for r in rows[: tpl.PROJECTS_CHAT_ROWS]]
        rest = len(rows) - len(per)
        more = f"\nЕщё проектов: {rest}, полный список в таблице отчёта." if rest > 0 else ""
        out["text_ru"] = head + " " + MEMBER_NO_FILE_TEXT + ("\n" + "\n".join(per) if per else "") + more
        return bi._safe(out)
    from app.web.routers.project_reports import SHARED_PREFIX, shared_report_url

    base = shared_reports_dir() / f"{SHARED_PREFIX}{_stamp()}"
    xlsx_path, html_path = base.with_suffix(".xlsx"), base.with_suffix(".html")
    gpt_rows = [{**n, "project_id": r["project_id"]} for r in found for n in r["gpt_nodes"]]
    sheets = [(tpl.PROJECTS_SHEET, columns, with_no_data(rows, ("spent_usd",)))]
    if gpt_rows:
        gpt_columns = tpl.PROJECTS_GPT_COLUMNS if show_spent else _without_spent(tpl.PROJECTS_GPT_COLUMNS)
        sheets.append((tpl.GPT_SHEET, gpt_columns, with_no_data(gpt_rows, ("cost_usd",))))
    write_workbook(xlsx_path, sheets, pairs)
    html_path.write_text(
        report_html.render_projects_html(
            bi._safe(
                {
                    "rows": rows,
                    "summary": summary,
                    "stamp": datetime.now().strftime("%d.%m.%Y %H:%M"),
                    # Участнику — без колонок трат, как в xlsx и чате.
                    "show_costs": show_spent,
                }
            )
        ),
        encoding="utf-8",
    )
    html_url, xlsx_url = shared_report_url(html_path), shared_report_url(xlsx_path)
    primary = xlsx_path if fmt == "xlsx" else html_path
    out.update(
        path=str(primary),
        download_url=shared_report_url(primary),
        file_name=primary.name,
        rows=len(rows),
        file_format=primary.suffix.lstrip("."),
        html_path=str(html_path),
        xlsx_path=str(xlsx_path),
        downloads=[{"url": html_url, "file_format": "html"}, {"url": xlsx_url, "file_format": "xlsx"}],
    )
    link = f"[Скачать отчёт (html)]({html_url}) · [Скачать отчёт (xlsx)]({xlsx_url})"
    shown = with_no_data(rows[: tpl.PROJECTS_CHAT_ROWS], ("spent_usd",))
    table = markdown_table(shown, chat_columns)
    rest = len(rows) - len(shown)
    tail = f"\n\nЕщё проектов: {rest} (полный список в файле отчёта)." if rest > 0 else ""
    gpt_lines = [
        f"- #{r['project_id']}: {clip(r['gpt'], 160)}" for r in shown if r.get("gpt") not in (None, "", "нет")
    ]
    gpt_text_md = ("\n\n" + tpl.GPT_SECTION_TITLE + ":\n" + "\n".join(gpt_lines)) if gpt_lines else ""
    out["markdown"] = f"## {tpl.PROJECTS_TITLE}\n\n{head}\n\n{table}{tail}{gpt_text_md}"
    # По строке на проект: без них «сравни проекты 5 и 7» видел только общие суммы.
    # Строка на каждый проект отчёта: при 10 строках из 14 ответ говорил «подробности о 10 из 14».
    # Участнику — без трат: суммы в $ видит только админ.
    per = [_project_line(r, show_spent=show_spent) for r in rows[: tpl.PROJECTS_CHAT_ROWS]]
    more = f"\nЕщё проектов: {rest}, полный список в файле отчёта." if rest > 0 else ""
    out["text_ru"] = head + f" Файл отчёта готов: {link}" + ("\n" + "\n".join(per) if per else "") + more
    return bi._safe(out)
