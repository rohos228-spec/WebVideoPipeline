"""Промты картинок (img_pr) из чата: не платить за пустой вход и честно назвать пустой выход.

Шаг img_pr строит промт кадра из данных сцены (место, действие, свет, шоты), которые
пишет нода «Работа с GPT». Если она не отработала (очередь пропустила её как готовую),
у кадров есть только закадр, а модель на каждый кадр отвечает «нет исходных данных для
заполнения». Студия ловит такой промт, только когда в нём три точные фразы-заглушки;
остальные варианты («данных о фоне нет», «исходные данные отсутствуют») проходят как
готовые. Здесь:

- перед запуском img_pr (шагом, цепочкой автопродвижения или очередью) читаем кадры
  проекта и отказываем, если у кадров, которым нужен промт, нет ни одного поля сцены;
  нода сцены, которая ещё ждёт, но которую очередь уже проскочила, в счёт не идёт;
- после img_pr считаем пустые промты устойчивой проверкой и не называем шаг успешным.
"""

from __future__ import annotations

import re
from collections import deque
from collections.abc import Iterable
from types import SimpleNamespace
from typing import Any

from loguru import logger

STEP_CODE = "img_pr"
IMG_PR_NODE_TYPE = "image_prompts"
# Ноды, которые пишут поля сцены кадра (кроме «Работы с GPT»): веер scene_design.
_SCENE_NODE_TYPES = ("scene_design", "sd_agent", "sd_assemble")
# Нода не будет запущена заново: очередь и автопродвижение идут мимо неё.
_NOT_RERUN = ("done", "skipped")
# Нода уже бралась в работу (шла, ждёт решения, готова или упала).
_STARTED = ("queued", "running", "waiting_hitl", "waiting_approval", "done", "failed")
# Поля, которые есть у кадра и без постановки сцены: персонажей пишет реестр героев.
_NOT_SCENE_KEYS = frozenset({"characters", "персонажи"})
_SHOW_NUMBERS = 12

# --- вход: есть ли у кадра данные сцены -------------------------------------------------


def scene_fields(frame: Any) -> list[str]:
    """Поля сцены кадра, которые шаг img_pr отдаст модели (meaning и attrs постановки)."""
    from app.services.db_frames_context import _pick_attrs

    out = ["meaning"] if str(getattr(frame, "meaning", None) or "").strip() else []
    picked = _pick_attrs(getattr(frame, "attrs", None))
    return out + sorted(k for k in picked if k not in _NOT_SCENE_KEYS)


def needs_prompt(frame: Any) -> bool:
    """Кадр, которому img_pr будет писать промт (тот же отбор, что у шага).

    На этой линии шаг берёт кадры с закадром и шоты многокадровой сцены
    (``xlsx_step_runners._is_real_shot``) без промта или с одними заглушками.
    """
    from app.orchestrator.steps.generate_image_prompts import _image_prompt_missing, _is_real_shot

    return _is_real_shot(frame) and _image_prompt_missing(getattr(frame, "image_prompt", None))


def frames_without_scene(frames: Iterable[Any]) -> tuple[list[int], int]:
    """Номера кадров без данных сцены среди тех, кому нужен промт, и сколько всего нужно."""
    need = [fr for fr in frames if needs_prompt(fr)]
    empty = [int(getattr(fr, "number", 0) or 0) for fr in need if not scene_fields(fr)]
    return empty, len(need)


def _snapshot(fr: Any) -> SimpleNamespace:
    return SimpleNamespace(
        number=fr.number,
        uuid=fr.uuid,
        voiceover_text=fr.voiceover_text,
        meaning=fr.meaning,
        image_prompt=fr.image_prompt,
        attrs=dict(fr.attrs) if isinstance(fr.attrs, dict) else {},
    )


async def load_live_frames(pid: int) -> list[SimpleNamespace] | None:
    """Живые (не архивные) кадры проекта из единой базы (project.db на этой линии нет).

    None - кадры прочитать не удалось: решать по ним нельзя.
    """
    from sqlalchemy import select

    from app.models import Frame, Project
    from app.services.action_layer import builtin_insight as bi
    from app.services.db_v2 import frame_is_archived

    xlsx: Any = None
    try:
        async with bi._sessions()() as s:
            rows = (await s.scalars(select(Frame).where(Frame.project_id == pid))).all()
            frames = [_snapshot(fr) for fr in rows if not frame_is_archived(fr)]
            if not rows:
                project = await s.get(Project, pid)
                xlsx = project.data_dir / "project.xlsx" if project is not None else None
    except Exception:  # noqa: BLE001 - база не читается: проверку не делаем
        return None
    if not rows and xlsx is not None:
        # В базе у проекта нет ни одной строки кадра (и архивной тоже), а project.xlsx есть:
        # так выглядит проект, перенесённый с линии, где кадры жили в project.db. Доска
        # монтажа и шаги заведут их из xlsx при первом обращении, и промты картинок пойдут
        # по ним же. Решаем по тем же кадрам. Все кадры в архиве - это решение человека,
        # xlsx их не воскрешает: живых кадров нет.
        frames = await xlsx_frames_preview(pid, xlsx) or []
    return sorted(frames, key=lambda f: int(f.number or 0))


# Превью кадров из xlsx на процесс: (проект, путь, mtime, размер) → кадры. Файл не менялся -
# повторная проверка в том же ходе чата (и в следующих) xlsx заново не разбирает.
_PREVIEW_CACHE: dict[tuple[int, str, int, int], tuple[SimpleNamespace, ...] | None] = {}
_PREVIEW_CACHE_MAX = 32
_PREVIEW_LOCKS: dict[tuple[int, str, int, int], Any] = {}


def _preview_key(pid: int, xlsx: Any) -> tuple[int, str, int, int] | None:
    try:
        if not xlsx.is_file():
            return None
        st = xlsx.stat()
    except OSError:
        return None
    return int(pid), str(xlsx), int(getattr(st, "st_mtime_ns", 0) or 0), int(getattr(st, "st_size", 0) or 0)


def clear_preview_cache() -> None:
    _PREVIEW_CACHE.clear()
    _PREVIEW_LOCKS.clear()


async def xlsx_frames_preview(pid: int, xlsx: Any = None) -> list[SimpleNamespace] | None:
    """Кадры, которые появятся в базе после импорта project.xlsx (как у доски монтажа).

    Импорт идёт в отдельной сессии и всегда откатывается: в базу ничего не пишется,
    событий не шлётся. Разбор xlsx (openpyxl) идёт в отдельном потоке со своим циклом
    событий, чтобы не держать чат; результат кэшируется по (проект, путь, mtime, размер).
    None - xlsx нет, он не читается или импорт не удался.
    """
    import asyncio

    if xlsx is None:
        xlsx = await _project_xlsx(pid)
        if xlsx is None:
            return None
    key = _preview_key(pid, xlsx)
    if key is None:
        return None
    if key in _PREVIEW_CACHE:
        cached = _PREVIEW_CACHE[key]
        return list(cached) if cached is not None else None
    lock = _PREVIEW_LOCKS.setdefault(key, asyncio.Lock())
    async with lock:
        if key not in _PREVIEW_CACHE:
            frames = await asyncio.to_thread(_preview_in_thread, pid)
            while len(_PREVIEW_CACHE) >= _PREVIEW_CACHE_MAX:
                _PREVIEW_CACHE.pop(next(iter(_PREVIEW_CACHE)))
            _PREVIEW_CACHE[key] = tuple(frames) if frames is not None else None
    _PREVIEW_LOCKS.pop(key, None)
    cached = _PREVIEW_CACHE.get(key)
    return list(cached) if cached is not None else None


async def _project_xlsx(pid: int) -> Any:
    from app.models import Project
    from app.services.action_layer import builtin_insight as bi

    try:
        async with bi._sessions()() as s:
            project = await s.get(Project, pid)
            return project.data_dir / "project.xlsx" if project is not None else None
    except Exception:  # noqa: BLE001 - база не читается: проверку не делаем
        return None


def _preview_in_thread(pid: int) -> list[SimpleNamespace] | None:
    """Поток превью: свой цикл событий, логи импорта помечены превью и опущены до DEBUG."""
    import asyncio

    from app.services.xlsx_preview_log import preview_logs

    with preview_logs():
        return asyncio.run(_preview_rolled_back(pid))


def _thread_sessions() -> tuple[Any, Any]:
    """Фабрика сессий для потока превью и движок, который потом закрыть (или None).

    Пул соединений Postgres привязан к главному циклу: в потоке свой движок без пула.
    SQLite и так без пула (NullPool), подмена из тестов - как есть.
    """
    from app.services.action_layer import builtin_insight as bi

    factory = bi._sessions()
    try:
        from app import db
        from app.settings import settings
    except Exception:  # noqa: BLE001
        return factory, None
    if factory is not db.SessionLocal or not settings.is_postgres:
        return factory, None
    from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
    from sqlalchemy.pool import NullPool

    engine = create_async_engine(settings.db_url, future=True, poolclass=NullPool)
    return async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession), engine


async def _preview_rolled_back(pid: int) -> list[SimpleNamespace] | None:
    from sqlalchemy import select

    from app.models import Frame, Project
    from app.services.db_v2 import frame_is_archived

    factory, engine = _thread_sessions()
    try:
        async with factory() as s:
            try:
                project = await s.get(Project, pid)
                if project is None:
                    return None
                xlsx = project.data_dir / "project.xlsx"
                if not xlsx.is_file():
                    return None
                if not await _import_xlsx_rolled_back(s, project, xlsx):
                    return None
                rows = (await s.scalars(select(Frame).where(Frame.project_id == pid))).all()
                return [_snapshot(fr) for fr in rows if not frame_is_archived(fr)]
            finally:
                await s.rollback()
    except Exception:  # noqa: BLE001 - xlsx или база не читаются: проверку не делаем
        logger.opt(exception=True).debug("img_pr guard: xlsx frames preview failed for #{}", pid)
        return None
    finally:
        if engine is not None:
            await engine.dispose()


async def _import_xlsx_rolled_back(session: Any, project: Any, xlsx: Any) -> bool:
    """Ядро ``chatgpt_xlsx.sync_project_xlsx`` без пометки ноды упавшей при сбое импорта."""
    from openpyxl import load_workbook

    from app.services.xlsx_sync import reload_from_xlsx
    from app.services.xlsx_v8_import import has_v8_plan_sheet, import_v8_xlsx
    from app.services.xlsx_versioning import validate_xlsx

    if validate_xlsx(xlsx) is not None:
        return False
    wb = load_workbook(filename=str(xlsx), data_only=True, read_only=True)
    try:
        is_v8 = has_v8_plan_sheet(wb)
    finally:
        wb.close()
    ok = False
    if is_v8:
        info = await import_v8_xlsx(session, project, xlsx, keep_fields=True, update_frames_voiceover=True)
        ok = not (isinstance(info, dict) and info.get("error"))
    try:
        await reload_from_xlsx(session, project, xlsx)
        ok = True
    except Exception:  # noqa: BLE001 - v7-лист не обязателен для v8-файла
        pass
    return ok


# --- выход: пустой ли промт ---------------------------------------------------------------

# «Метка: текст» внутри промта (Фон: …, Действие: …, Scene sense (visible): …).
_LABEL_RE = re.compile(r"(?:^|(?<=[.!?\n]))\s*([A-Za-zА-ЯЁ][^:.;!?\n]{0,60}):")
_CLAUSE_END_RE = re.compile(r"[;.!?\n]")
_STYLE_RE = re.compile(r"\bSTYLE\s*:")
# Заглушки модели вместо содержания: «нет исходных данных», «исходные данные о фоне
# отсутствуют», «данных о фоне кадра нет», «источник и направление не заданы».
_PLACEHOLDER_RES = tuple(
    re.compile(p, re.IGNORECASE)
    for p in (
        r"\bнет\s+(?:\S+\s+){0,2}исходн\w*\s+(?:\S+\s+){0,2}данн",
        r"\bисходн\w*\s+(?:\S+\s+){0,3}данн\w*.{0,90}?"
        r"(?:отсутству|не\s+предоставлен|не\s+задан|не\s+указан|недоступн|не\s+определ|\bнет\b)",
        r"\bисходн\w*\s+[\w/, ]{1,60}?\s(?:отсутству|не\s+предоставлен|не\s+задан|не\s+указан)",
        # Текст-рыба шаблона вместо сцены.
        r"\blorem\s+ipsum\b",
        r"\bданн\w*\s+(?:о|об|для|по)\b.{0,90}?\bнет\b",
        # Отсутствие именно ВХОДНЫХ сведений (данных/информации/описания), а не предмета сцены:
        # «тень отсутствует», «людей нет», «источник не указан, мягкий» — это содержание.
        r"\b(?:данн\w*|информаци\w*|описани\w*|сведени\w*)\b.{0,90}?"
        r"(?:отсутству|не\s+задан|не\s+указан|не\s+предоставлен|\bнет\b)",
        r"\b(?:нет|без)\s+(?:\S+\s+){0,2}(?:данн|информаци|описани|сведени)",
        # «Недостаточно исходных данных для заполнения» — тот же отказ другими словами.
        r"\bнедостаточно\s+(?:\S+\s+){0,2}(?:данн|информаци|сведени)",
        # Английские формы: «No source data available…», «Insufficient input data».
        r"\b(?:no|insufficient|missing|not\s+enough)\s+(?:\S+\s+){0,2}"
        r"(?:source|input|scene)\s+(?:data|information|details)\b",
        r"\b(?:source|input)\s+(?:data|information)\b.{0,60}?"
        r"\b(?:missing|unavailable|absent|not\s+(?:available|provided|specified|given))",
        r"\bno\s+(?:data|information)\s+(?:is\s+)?(?:available|provided|given)\b",
        # «Свет: источник, направление и характер света не указаны[; не добавлять …]»: параметры
        # раздела не заданы и дальше нет описания. «источник не указан, мягкий рассеянный» —
        # за «не указан» идёт описание, это содержание. («отсутствует» здесь не берём:
        # «тень отсутствует» — предмет сцены.)
        r"\b(?:не\s+задан\w*|не\s+указан\w*|не\s+предоставлен\w*)\s*"
        r"(?:$|,\s*(?:не\b|без\b|сохран|остав))",
        # Раздел из одного слова-отсутствия: «Свет: не задан.», «Фон: отсутствует.»
        r"^\s*(?:не\s+задан\w*|не\s+указан\w*|не\s+предоставлен\w*|отсутству\w*|нет)\s*$",
    )
)
# Значение раздела вместо сцены: прочерк, «н/д», «неизвестно» («Фон: —. Действие: н/д.»).
# Только всё значение раздела с меткой: «неизвестный путник идёт» — содержание.
# Тоже целиком: «не определено», «пусто», «ничего», «см. выше», «not specified».
_EMPTY_VALUE_RE = re.compile(
    r"^\s*(?:[-—–…]+|н/д|n/?a|неизвестн\w*|unknown|none|tbd|\?+|пусто|ничего|nothing|empty|"
    r"не\s+определен\w*|не\s+определ[её]н\w*|undefined|unspecified|"
    r"not\s+(?:specified|provided|given|defined|set|available|known)|"
    r"см\.?\s*(?:выше|ниже|ранее)|как\s+выше|аналогично(?:\s+выше)?|see\s+above|same\s+as\s+above)\s*$",
    re.IGNORECASE,
)
# Слот шаблона в скобках вместо текста: «[фон]», «[описание действия]», «{scene}», «<background>».
_STUB_SLOT = (
    r"\s*(?:описани\w*\s+)?(?:фон\w*|действи\w*|свет\w*|мест\w*|локаци\w*|сцен\w*|персонаж\w*|"
    r"кадр\w*|детал\w*|текст\w*|background|action|light\w*|scene|setting|location|characters?|"
    r"description|text|placeholder|insert[^\]\}>\n]{0,30}|вставить[^\]\}>\n]{0,30}|todo|tbd|…|\.\.\.)\s*"
)
_STUB_RE = re.compile(rf"\[{_STUB_SLOT}\]|\{{{_STUB_SLOT}\}}|<{_STUB_SLOT}>", re.IGNORECASE)
# Разделы сцены, которые бывают и после STYLE: тогда промт без текста до STYLE не пуст.
_SCENE_LABEL_RE = re.compile(
    r"(?:^|[.!?\n]\s*)(?:фон|действие|сцена|место|персонаж\w*|композиц\w*|"
    r"background|scene|action|setting|characters?|composition|subject)\s*:",
    re.IGNORECASE,
)
# Разделы не про сцену: референс героя, запреты, стиль.
_NOT_SCENE_LABEL_RE = re.compile(
    r"референс|reference|детал|запрет|negative|style|палитр|palette", re.IGNORECASE
)
# Раздел-запрет («не добавлять …») ни содержанием, ни заглушкой не считаем.
_NEGATIVE_RE = re.compile(r"^\s*(?:не|без|no|do not|never)\b", re.IGNORECASE)


def _is_placeholder(text: str) -> bool:
    flat = " ".join(str(text or "").split())
    return any(r.search(flat) for r in _PLACEHOLDER_RES)


# «нет исходных данных о фоне, но лиса бежит по снегу»: после «но/однако/зато» идёт сцена.
# С запятой и без: «…данных о фоне но лиса бежит» — то же самое.
_CONTRAST_RE = re.compile(r"(?:,\s*|\s+)(?:но|однако|зато|but|however)\s+", re.IGNORECASE)
# После «но» указание про стиль / «сохранить нейтральный фон» / «оставить пустым» —
# не сцена: «нет исходных данных, но сохранить нейтральный фон» остаётся заглушкой.
# «но STYLE LOCK обязательный», «но кадр должен быть атмосферным» — указание модели, не сцена.
_NOT_SCENE_REST_RE = re.compile(
    r"\bстил|\bstyle\b|^(?:сохран|остав|держ|следова|соблюд|использ|придерж)|нейтральн|\bпуст|"
    r"^(?:кадр\w*|изображени\w*|картинк\w*|иллюстраци\w*|промт\w*|сцен\w*|атмосфер\w*|frame|image|"
    r"illustration|prompt|scene)\s+(?:\S+\s+)?(?:долж|нуж|обязан|следует|must|should|needs?\b)",
    re.IGNORECASE,
)


def _content_after_contrast(clause: str) -> bool:
    parts = _CONTRAST_RE.split(clause, maxsplit=1)
    if len(parts) < 2:
        return False
    rest = parts[1].strip()
    return (
        bool(rest)
        and not _is_placeholder(rest)
        and not _NEGATIVE_RE.match(rest)
        and not _NOT_SCENE_REST_RE.search(rest)
    )


def _sections(body: str) -> list[tuple[str, str]]:
    marks = list(_LABEL_RE.finditer(body))
    if not marks:
        return [("", body.strip())] if body.strip() else []
    out = []
    head = body[: marks[0].start()].strip()
    if head:
        out.append(("", head))
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(body)
        out.append((m.group(1).strip(), body[m.end() : end].strip()))
    return out


def prompt_is_hollow(prompt: str | None) -> bool:
    """Промт без содержания сцены: пустой, или все его разделы сцены - заглушки.

    Стиль (STYLE: … и дальше) и разделы референса/запретов в счёт не идут: они есть у
    любого промта. Раздел-заглушка - тот, что начинается со «нет исходных данных» и
    похожих фраз; раздел с содержанием - любой другой непустой.
    """
    from app.generation_options import is_scene_prompt_hollow, strip_prompt_id_lines

    text = strip_prompt_id_lines(str(prompt or "")).strip()
    if not text:
        return True
    if is_scene_prompt_hollow(text):
        return True
    body, *after_style = _STYLE_RE.split(text, maxsplit=1)
    # «[фон]» и подобные слоты шаблона — не текст сцены.
    body = _STUB_RE.sub(" ", body)
    if not re.search(r"\w", body):
        # До STYLE ни слова: промт — один стилевой замок, если и после него нет раздела сцены.
        return not _SCENE_LABEL_RE.search(after_style[0] if after_style else "")
    placeholders = content = 0
    for label, part in _sections(body):
        if _is_placeholder(label):
            placeholders += 1
            continue
        if _NOT_SCENE_LABEL_RE.search(label):
            continue
        clauses = [c for c in _CLAUSE_END_RE.split(part) if c.strip()]
        # «Фон: —. Тени длинные» — решает первое предложение после прочерка.
        filled = [c for c in clauses if not _EMPTY_VALUE_RE.match(c)] if label else clauses
        first = filled[0] if filled else ""
        if label and (not filled or _EMPTY_VALUE_RE.match(part.strip().rstrip(".;!?"))):
            # «Фон: .», «Фон: —», «Фон: н/д», «Фон: см. выше» — раздел без сцены.
            placeholders += 1
        elif _is_placeholder(first) and _content_after_contrast(first):
            content += 1
        elif _is_placeholder(first):
            placeholders += 1
        elif first.strip() and not _NEGATIVE_RE.match(first):
            content += 1
    return placeholders > 0 and content == 0


def hollow_frames(frames: Iterable[Any]) -> tuple[list[int], int]:
    """Номера кадров с пустым промтом среди кадров шага img_pr и сколько их всего."""
    from app.orchestrator.steps.generate_image_prompts import _is_real_shot

    parents = [fr for fr in frames if _is_real_shot(fr)]
    hollow = [
        int(getattr(fr, "number", 0) or 0)
        for fr in parents
        if prompt_is_hollow(getattr(fr, "image_prompt", None))
    ]
    return hollow, len(parents)


# --- граф: дойдёт ли запуск до img_pr мимо ноды сцены --------------------------------------


def _graph(meta: Any) -> tuple[dict[str, dict], dict[str, list[str]], dict[str, list[str]]]:
    from app.services.canvas_graph import canvas_graph_from_meta

    cg = canvas_graph_from_meta(meta if isinstance(meta, dict) else {}) or {}
    nodes = {str(n["id"]): n for n in cg.get("nodes") or [] if isinstance(n, dict) and n.get("id")}
    outs: dict[str, list[str]] = {}
    ins: dict[str, list[str]] = {}
    for e in cg.get("edges") or []:
        if not isinstance(e, dict):
            continue
        s, t = str(e.get("source") or "").strip(), str(e.get("target") or "").strip()
        if s and t:
            outs.setdefault(s, []).append(t)
            ins.setdefault(t, []).append(s)
    return nodes, outs, ins


def _type(node: dict | None) -> str:
    from app.services.excel_gpt_node import effective_node_type

    return effective_node_type(node or {})


def fills_scene(key: str, node: dict | None) -> bool:
    """Нода пишет поля сцены кадра: «Работа с GPT» (кроме код-ноды отчёта) или веер сцен."""
    from app.services.excel_gpt_node import is_excel_gpt_node_type

    if key.endswith("_fw_report") or "_fw_report_" in key:
        return False
    typ = _type(node)
    raw = str((node or {}).get("type") or "")
    return typ in _SCENE_NODE_TYPES or is_excel_gpt_node_type(typ) or is_excel_gpt_node_type(raw)


def _statuses(node_runs: Iterable[dict] | None) -> dict[str, str]:
    return {
        str(nr.get("node_key")): str(nr.get("status") or "")
        for nr in node_runs or []
        if isinstance(nr, dict) and nr.get("node_key")
    }


def _img_pr_keys(nodes: dict[str, dict]) -> list[str]:
    return [k for k, n in nodes.items() if _type(n) == IMG_PR_NODE_TYPE]


def step_node_key(meta: Any, step_code: str) -> str | None:
    """Нода графа, с которой стартует шаг ``step_code`` (первая такого типа)."""
    from app.orchestrator.node_registry import STEP_CODE_TO_NODE_TYPE

    want = STEP_CODE_TO_NODE_TYPE.get(step_code)
    if not want:
        return None
    nodes, _outs, _ins = _graph(meta)
    return next((k for k, n in nodes.items() if _type(n) == want), None)


def scene_nodes_before_img_pr(meta: Any) -> list[str]:
    """Ноды, которые пишут сцены для промтов картинок: ближние к ним ноды сцены по стрелкам.

    Дальше ближней ноды вверх не идём: «Работа с GPT» выше неё (каркас, разметка героев)
    сцен не пишет - в #10/#12/#14 первая такая нода готова, а у кадров нет ни одного поля
    сцены. В порядке от ближней.
    """
    nodes, _outs, ins = _graph(meta)
    found: list[str] = []
    seen: set[str] = set()
    queue: deque[str] = deque(k for key in _img_pr_keys(nodes) for k in ins.get(key, []))
    while queue:
        key = queue.popleft()
        if key in seen:
            continue
        seen.add(key)
        if fills_scene(key, nodes.get(key)):
            found.append(key)
            continue
        queue.extend(ins.get(key, []))
    return found


def _descendants(outs: dict[str, list[str]], key: str) -> set[str]:
    seen: set[str] = set()
    queue: deque[str] = deque(outs.get(key, []))
    while queue:
        nxt = queue.popleft()
        if nxt not in seen:
            seen.add(nxt)
            queue.extend(outs.get(nxt, []))
    return seen


def _stale_gpt_key(meta: Any, key: str, node: dict | None) -> bool:
    """Студия закроет «Работу с GPT» ``key`` без вызова модели.

    Так было в #12: в meta осталась активной уже готовая первая нода, очередь вошла в слот
    второй, и шаг записал в журнал «node=n_excel_gpt_1 already in completed_keys - skip GPT
    (leftover worker tick)»; вторая нода осталась pending, а очередь пошла к промтам картинок.
    """
    from app.services.excel_gpt_node import excel_gpt_already_completed_no_force, is_excel_gpt_node_type

    if not (
        is_excel_gpt_node_type(_type(node)) or is_excel_gpt_node_type(str((node or {}).get("type") or ""))
    ):
        return False
    m = meta if isinstance(meta, dict) else {}
    active = str(m.get("active_excel_gpt_node_key") or "").strip()
    return bool(active) and active != key and excel_gpt_already_completed_no_force(SimpleNamespace(meta=m))


def jumped_scene_nodes(
    meta: Any, node_runs: Iterable[dict] | None, *, explicit_gpt_start: bool = False
) -> list[str]:
    """Ноды сцены перед промтами картинок, которые ещё не готовы, но очередь их проскочит.

    Признаки: после ноды по стрелкам уже шли шаги (готовы, упали или идут), а она так и
    ждёт; или Студия закроет её без вызова модели (см. ``_stale_gpt_key``). Явный запуск
    ноды «Работа с GPT» из чата снимает залипший ключ - тогда второй признак не считаем.
    """
    nodes, outs, _ins = _graph(meta)
    status = _statuses(node_runs)
    out: list[str] = []
    for key in scene_nodes_before_img_pr(meta):
        if status.get(key) in _NOT_RERUN:
            continue
        ran_after = any(status.get(k) in _STARTED for k in _descendants(outs, key))
        if ran_after or (not explicit_gpt_start and _stale_gpt_key(meta, key, nodes.get(key))):
            out.append(key)
    return out


def _will_not_fill(
    meta: Any, node_runs: Iterable[dict] | None, *, explicit_gpt_start: bool = False
) -> set[str]:
    status = _statuses(node_runs)
    idle = {k for k in scene_nodes_before_img_pr(meta) if status.get(k) in _NOT_RERUN}
    return idle | set(jumped_scene_nodes(meta, node_runs, explicit_gpt_start=explicit_gpt_start))


def launch_skips_scene(
    meta: Any, step_code: str, node_key: str | None, node_runs: Iterable[dict] | None, *, auto_chain: bool
) -> bool:
    """Запуск шага дойдёт до промтов картинок без ноды сцены перед ними.

    Сам шаг img_pr - да. С автопродвижением: ранние шаги (план, сценарий, разбивка)
    сбрасывают ноды ниже, и ноды сцены пойдут заново; нода, которая пишет сцены, как старт
    сама их заполнит; иначе смотрим, есть ли по стрелкам до промтов нода сцены, которая
    действительно запустится.
    """
    from app.services import launch_risks as lr

    if step_code == STEP_CODE:
        return True
    if not auto_chain or step_code in lr.EARLY_STEPS:
        return False
    start = node_key or step_node_key(meta, step_code)
    if not start:
        return False
    if start in scene_nodes_before_img_pr(meta):
        return False
    return chain_skips_scene(meta, start, node_runs, explicit_gpt_start=step_code == "excel_gpt")


def chain_skips_scene(
    meta: Any, start_key: str, node_runs: Iterable[dict] | None, *, explicit_gpt_start: bool = False
) -> bool:
    """Цепочка от ``start_key`` по стрелкам дойдёт до промтов картинок, не запустив ни одной
    ноды сцены (готовые и пропущенные воркер не перезапускает, проскакиваемые - тоже)."""
    nodes, outs, _ins = _graph(meta)
    fillers = set(scene_nodes_before_img_pr(meta))
    idle = _will_not_fill(meta, node_runs, explicit_gpt_start=explicit_gpt_start)
    seen: set[str] = set()
    queue: deque[str] = deque(outs.get(start_key, []))
    while queue:
        key = queue.popleft()
        if key in seen:
            continue
        seen.add(key)
        if _type(nodes.get(key)) == IMG_PR_NODE_TYPE:
            return True
        if key in fillers and key not in idle:
            continue
        queue.extend(outs.get(key, []))
    return False


def queue_skips_scene(
    meta: Any,
    node_runs: Iterable[dict] | None,
    *,
    mode: str = "full",
    target_key: str | None = None,
    target_type: str | None = None,
) -> bool:
    """Очередь дойдёт до промтов картинок, а ни одна нода сцены перед ними не запустится.

    Очередь идёт с текущего места проекта: готовые и пропущенные ноды не повторяет и
    проскакивает ноды из ``jumped_scene_nodes``. Цель until_node выше промтов картинок -
    очередь до них не дойдёт.
    """
    nodes, outs, _ins = _graph(meta)
    img_keys = _img_pr_keys(nodes)
    status = _statuses(node_runs)
    if not img_keys or all(status.get(k) in _NOT_RERUN for k in img_keys):
        return False
    if mode == "until_node" and not _target_at_or_after(nodes, outs, img_keys, target_key, target_type):
        return False
    idle = _will_not_fill(meta, node_runs)
    return all(k in idle for k in scene_nodes_before_img_pr(meta))


# Статусы, с которых «Продолжить» без автопродвижения следующим шагом запустит промты картинок:
# сам шаг (пауза или сбой на нём) и готовность прямо перед ним.
_CONTINUE_INTO_IMG_PR = frozenset({"generating_image_prompts", "enrich_5_ready"})


def _continue_place(meta: Any, status: str | None) -> str:
    """Откуда «Продолжить» поведёт проект: с паузы - шаг до паузы или шаг, на котором упал."""
    if str(status or "") != "paused":
        return str(status or "")
    m = meta if isinstance(meta, dict) else {}
    fs = m.get("step_failure") if isinstance(m.get("step_failure"), dict) else {}
    return str(m.get("paused_from_status") or fs.get("last_running") or "")


def _into_img_pr(enrich_slots: Any) -> frozenset[str]:
    """Места, откуда следующий шаг - промты картинок: сам шаг и готовность последнего
    включённого слота «Доп работа с EXCEL» (сдача шага после ошибок пишет её в
    ``paused_from_status``)."""
    try:
        n = int(enrich_slots or 0)
    except (TypeError, ValueError):
        n = 0
    return _CONTINUE_INTO_IMG_PR | ({f"enrich_{n}_ready"} if 1 <= n <= 5 else set())


def continue_skips_scene(
    meta: Any,
    node_runs: Iterable[dict] | None,
    *,
    status: str | None,
    auto_mode: bool,
    enrich_slots: int | None = None,
) -> bool:
    """«Продолжить конвейер» дойдёт до промтов картинок, а ни одна нода сцены перед ними не
    запустится.

    Так было в #10: «Работа с GPT 2» ждёт, промты картинок уже падали после неё, и
    продолжение идёт с места проекта - снова к промтам. С автопродвижением цепочка дойдёт до
    промтов сама; без него - только если промты и есть следующий шаг.
    """
    if not queue_skips_scene(meta, node_runs):
        return False
    if auto_mode or jumped_scene_nodes(meta, node_runs):
        return True
    place = _continue_place(meta, status)
    return place in _into_img_pr(enrich_slots) or _ready_slot_leads_to_img_pr(meta, place)


_ENRICH_READY_RE = re.compile(r"^enrich_([1-5])_ready$")


def _explicit_slot(node: dict) -> int | None:
    """Номер слота, записанный в ноде (slotIndex или старый тип enrich_N); без номера - None.

    ``slot_index_from_node`` считает ноду без номера первым слотом: для проверки продолжения
    это догадка, по ней не блокируем.
    """
    from app.services.excel_gpt_node import MAX_EXCEL_GPT_SLOTS, legacy_enrich_slot_from_type

    data = node.get("data") if isinstance(node.get("data"), dict) else {}
    raw = data.get("slotIndex")
    if isinstance(raw, int) and not isinstance(raw, bool) and 1 <= raw <= MAX_EXCEL_GPT_SLOTS:
        return raw
    return legacy_enrich_slot_from_type(str(node.get("type") or ""))


def _ready_slot_node(meta: Any, nodes: dict[str, dict], slot: int) -> str | None:
    """Нода «Работа с GPT», на чью готовность указывает ``enrich_N_ready``; None - не понять.

    Сначала нода, которую воркер закрепил (``active_excel_gpt_node_key``), если номер её слота
    совпадает. Иначе единственная нода с этим номером, и только когда на холсте нет
    «Работ с GPT» без номера: такая нода считается первым слотом, и какая из них готова,
    по холсту не сказать.
    """
    from app.services.excel_gpt_node import is_excel_gpt_node_type, sd_agent_marker

    gpt = {
        k: n
        for k, n in nodes.items()
        if (is_excel_gpt_node_type(_type(n)) or is_excel_gpt_node_type(str(n.get("type") or "")))
        and not sd_agent_marker(n)
        and not (isinstance(n.get("data"), dict) and n["data"].get("slotOverflow"))
    }
    m = meta if isinstance(meta, dict) else {}
    active = str(m.get("active_excel_gpt_node_key") or "").strip()
    if active in gpt and _explicit_slot(gpt[active]) == slot:
        return active
    if any(_explicit_slot(n) is None for n in gpt.values()):
        return None
    numbered = [k for k, n in gpt.items() if _explicit_slot(n) == slot]
    return numbered[0] if len(numbered) == 1 else None


def _ready_slot_leads_to_img_pr(meta: Any, place: str) -> bool:
    """С готовности «Работы с GPT» слота N следующим шагом идут промты картинок.

    Воркер с ``enrich_N_ready`` идёт по стрелкам холста, а не по числу слотов в настройках:
    слотов в настройках может быть три, а на холсте две «Работы с GPT», и от второй стрелка
    ведёт прямо в промты картинок, так что «Продолжить» с ``enrich_2_ready`` запустит именно их.
    Стрелки читаем только от той ноды, на которую указывает статус; если её не определить
    (слоты без номеров, два одинаковых номера), это правило не блокирует - остальные
    проверки остаются. Служебные ноды (проходные) по пути пропускаем, любая другая рабочая
    нода - не промты.
    """
    from app.orchestrator.graph.planner import is_passthrough_node_type

    m = _ENRICH_READY_RE.match(str(place or ""))
    if not m:
        return False
    nodes, outs, _ins = _graph(meta)
    start = _ready_slot_node(meta, nodes, int(m.group(1)))
    if start is None:
        return False
    seen: set[str] = set()
    queue: deque[str] = deque(outs.get(start, []))
    while queue:
        key = queue.popleft()
        if key in seen:
            continue
        seen.add(key)
        typ = _type(nodes.get(key))
        if typ == IMG_PR_NODE_TYPE:
            return True
        if is_passthrough_node_type(typ):
            queue.extend(outs.get(key, []))
    return False


def _target_at_or_after(nodes, outs, img_keys, target_key, target_type) -> bool:
    seen: set[str] = set()
    queue: deque[str] = deque(img_keys)
    while queue:
        key = queue.popleft()
        if key in seen:
            continue
        seen.add(key)
        if (target_key and key == target_key) or (
            not target_key and target_type and _type(nodes.get(key)) == target_type
        ):
            return True
        queue.extend(outs.get(key, []))
    return False


# --- тексты ----------------------------------------------------------------------------------


def numbers_ru(numbers: list[int]) -> str:
    shown = ", ".join(str(n) for n in numbers[:_SHOW_NUMBERS])
    rest = len(numbers) - _SHOW_NUMBERS
    return shown + (f" и ещё {rest}" if rest > 0 else "")


def _nodes_ru(meta: Any, keys: list[str]) -> str:
    from app.services.action_layer.builtin_group import node_label_ru

    return ", ".join(f"«{node_label_ru(meta, k)}» ({k})" for k in keys)


def _gpt_nodes_ru(meta: Any, keys: list[str]) -> str:
    """«Работа с GPT» «Сцены» (n_gpt2); без повтора, если подпись ноды уже «Работа с GPT …»."""
    from app.services.action_layer.builtin_group import node_label_ru

    named = _nodes_ru(meta, keys)
    if all(str(node_label_ru(meta, k) or "").startswith("Работа с GPT") for k in keys):
        return named
    return f"«Работа с GPT» {named}"


def _idle_fillers(meta: Any, node_runs: Iterable[dict] | None) -> tuple[list[str], list[str]]:
    """Ноды сцены, которые не отработают (пропущенные и проскакиваемые), и из них проскакиваемые."""
    keys = scene_nodes_before_img_pr(meta)
    status = _statuses(node_runs)
    jumped = jumped_scene_nodes(meta, node_runs)
    idle = [k for k in keys if status.get(k) == "skipped" or k in jumped]
    return idle, jumped


def _who_fills(meta: Any, keys: list[str]) -> str:
    """«пишет нода …» / «пишут ноды …» по списку нод сцены."""
    if not keys:
        return "пишет нода «Работа с GPT», которая заполняет сцены"
    head = "пишет нода" if len(keys) == 1 else "пишут ноды"
    return f"{head} {_gpt_nodes_ru(meta, keys)}"


def no_scene_short(meta: Any, node_runs: Iterable[dict] | None) -> str:
    """Коротко для списка «не поставлены в очередь»: «кадры без данных сцены, сначала запустите ноду …»."""
    from app.services.action_layer.builtin_group import node_label_ru

    idle, jumped = _idle_fillers(meta, node_runs)
    keys = jumped or idle or scene_nodes_before_img_pr(meta)
    if not keys:
        return "кадры без данных сцены, сначала запустите ноду «Работа с GPT», которая заполняет сцены"
    names = ", ".join(f"«{node_label_ru(meta, k)}»" for k in keys)
    return f"кадры без данных сцены, сначала запустите {'ноду' if len(keys) == 1 else 'ноды'} {names}"


def no_scene_text(pid: int, empty: list[int], total: int, meta: Any, node_runs: Iterable[dict] | None) -> str:
    idle, jumped = _idle_fillers(meta, node_runs)
    keys = idle or scene_nodes_before_img_pr(meta)
    text = (
        f"у проекта #{pid} кадры {numbers_ru(empty)} ({len(empty)} из {total}, которым "
        "нужен промт) без данных сцены: нет ни места, ни действия, ни света, ни шотов, "
        f"только закадр. Эти поля {_who_fills(meta, keys)}, но в этом проекте "
        "кадры остались незаполненными"
    )
    if jumped:
        one = len(jumped) == 1
        text += (
            f". {'Нода' if one else 'Ноды'} {_nodes_ru(meta, jumped)} ещё не "
            f"{'отработала' if one else 'отработали'}, но очередь и автопродвижение "
            f"проходят мимо {'неё' if one else 'них'} к промтам картинок"
        )
    text += ". Промты картинок из таких кадров выйдут пустыми («нет исходных данных»), деньги уйдут зря. "
    if jumped:
        return text + (
            f"Сначала запустите {'эту ноду' if len(jumped) == 1 else 'эти ноды'} "
            "«Работа с GPT» отдельно, потом промты картинок"
        )
    return text + "Сначала запустите ноду сцен, потом промты картинок"


def hollow_text(
    hollow: list[int], total: int, *, no_scene: bool, kept: list[int] | None = None, scene_nodes: str = ""
) -> str:
    """Итог «шаг формально готов, но промты пустые …» и что делать дальше.

    ``kept`` - кадры, чьи пустые промты Студия считает готовыми: повторный запуск шага
    (из чата он всегда продолжает, без стирания) их пропустит.
    """
    parts = [
        f"шаг формально готов, но промты пустые у кадров {numbers_ru(hollow)} "
        f"({len(hollow)} из {total}): вместо сцены модель написала «нет исходных данных»"
    ]
    if no_scene:
        who = scene_nodes or "«Работа с GPT»"
        parts.append(f"У этих кадров нет данных сцены: сначала нужна нода {who}, которая их заполняет")
    if kept:
        parts.append(
            f"Повторный запуск промтов картинок кадры {numbers_ru(kept)} не перепишет: "
            "Студия считает их промты готовыми. Сотрите у них промт картинки в разделе "
            "«База» и только потом запускайте промты картинок заново"
        )
    elif no_scene:
        parts.append("Потом запустите промты картинок заново")
    parts.append("Картинки по таким промтам делать не стоит")
    return ". ".join(parts)


# --- проверки для действий -------------------------------------------------------------------


async def scene_gap(pid: int, meta: Any, node_runs: Iterable[dict] | None) -> str | None:
    """Отказ, если кадрам, которым нужен промт, нечем его наполнить; None - можно."""
    frames = await load_live_frames(pid)
    if not frames:
        return None
    empty, total = frames_without_scene(frames)
    if not empty:
        return None
    return no_scene_text(pid, empty, total, meta, node_runs)


async def hollow_outcome(pid: int, meta: Any = None) -> str | None:
    """Текст «шаг формально готов, но промты пустые …» или None, если промты с содержанием."""
    from app.generation_options import image_prompt_missing

    frames = await load_live_frames(pid)
    if not frames:
        return None
    hollow, total = hollow_frames(frames)
    if not hollow:
        return None
    by_num = {int(fr.number or 0): fr for fr in frames}
    rows = [by_num[n] for n in hollow if n in by_num]
    no_scene = all(not scene_fields(fr) for fr in rows)
    kept = [int(fr.number or 0) for fr in rows if not image_prompt_missing(fr.image_prompt)]
    keys = scene_nodes_before_img_pr(meta) if meta is not None else []
    return hollow_text(
        hollow, total, no_scene=no_scene, kept=kept, scene_nodes=_gpt_nodes_ru(meta, keys) if keys else ""
    )
